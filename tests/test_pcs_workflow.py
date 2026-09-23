from __future__ import annotations

import sqlite3
import tempfile
import unittest
import zipfile
from contextlib import nullcontext
from datetime import date, datetime
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import MagicMock, patch
from xml.etree import ElementTree

from openpyxl import load_workbook

from wfmhub.config import ConfigError, _pcs_workbook_path, write_pcs_workbook
from wfmhub.cli import (
    _build_latest_pcs_report,
    _build_pcs_from_database,
    _pcs_menu,
    _sync_pcs_workbook,
    refresh,
)
from wfmhub.models import ModelSummary
from wfmhub.database import DatabaseConnection
from wfmhub.pcs_excel import inspect_pcs_tracker
from wfmhub.pcs_tracker import (
    PCS_TRACKER_FILENAME,
    PCS_TRACKER_VERSION,
    _as_date,
    _tracker_contract_version,
    build_pcs_live_tracker,
    ensure_pcs_tracker,
    latest_pcs_report,
    tracker_path,
)
from wfmhub.shared_feeds import (
    PCS_AGENT_SCORECARD_HEADERS,
    PCS_COACHING_HEADERS,
    PCS_DAILY_SCORECARD_HEADERS,
    PCS_FILTER_HEADERS,
    PCS_LOB_SCORECARD_HEADERS,
    PCS_RESULTS_HEADERS,
    pcs_dashboard_cache_rows,
    pcs_result_rows,
    _pcs_reporting_windows,
)


def _rows() -> dict[str, list[tuple[object, ...]]]:
    refreshed = datetime(2026, 9, 8, 18, 0)
    selection = "Current MTD|All|All|All"
    filters = [
        ("Current MTD", "All", "TL|All", "All", "AGENT|All|All", "All"),
        ("", "RSA NL", "TL|All", "TL 1", "AGENT|All|All", "Agent One [001]"),
        ("", "", "TL|RSA NL", "All", "AGENT|All|TL 1", "All"),
        ("", "", "TL|RSA NL", "TL 1", "AGENT|All|TL 1", "Agent One [001]"),
        ("", "", "", "", "AGENT|RSA NL|All", "All"),
        ("", "", "", "", "AGENT|RSA NL|All", "Agent One [001]"),
        ("", "", "", "", "AGENT|RSA NL|TL 1", "All"),
        ("", "", "", "", "AGENT|RSA NL|TL 1", "Agent One [001]"),
    ]
    lob = (
        f"KPI|{selection}", "KPI", 0, "All", 20, 4.3, 4.1, .2,
        .15, 3, date(2026, 9, 8), refreshed,
    )
    lob_detail = (
        f"LOB|{selection}|1", "LOB", 1, "RSA NL", 20, 4.3, 4.1, .2,
        .15, 3, date(2026, 9, 8), refreshed,
    )
    agent = (
        f"AGENT|{selection}|1", 1, "Agent One", "001", "RSA NL", "TL 1",
        4.3, 4.0, .3, .15, 10, 2, date(2026, 9, 8), refreshed,
    )
    daily = (
        f"DAILY|{selection}|1", 1, date(2026, 9, 8), 4.3, .15, 10, 2,
        date(2026, 9, 8), refreshed,
    )
    result = (
        "Current MTD", date(2026, 9, 1), date(2026, 9, 8), "AGENT",
        "RSA NL", "TL 1", "Agent One [001]", "001", "Agent One", "NL",
        4.3, .15, 10, 60, 9, 2, 8, 120, "OK", date(2026, 9, 8), refreshed,
    )
    coaching = (
        "COACH|Current MTD|All|1", 1, date(2026, 9, 8), "RSA NL", "TL 1",
        "Agent One", "001", 2, "High", "call-1", "Review the explanation",
        "key-1", date(2026, 9, 8), refreshed,
    )
    return {
        "filter_rows": filters,
        "lob_rows": [lob, lob_detail],
        "agent_rows": [agent],
        "daily_rows": [daily],
        "result_rows": [result],
        "coaching_rows": [coaching],
    }


class PCSWorkflowTests(unittest.TestCase):
    def test_dedicated_pcs_path_preserves_the_original_and_coaching(self):
        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            local = SimpleNamespace(reports=root / "Reports", feed=root / "Feed")
            original = ensure_pcs_tracker(local, **_rows())
            workbook = load_workbook(original)
            workbook["COACHING"]["L5"] = "key-shared"
            workbook["COACHING"]["N5"] = "In progress"
            workbook.save(original)
            workbook.close()
            original_bytes = original.read_bytes()

            shared_dir = root / "Synced WFM SharePoint"
            shared_dir.mkdir()
            shared = shared_dir / PCS_TRACKER_FILENAME
            config = SimpleNamespace(
                reports=local.reports, feed=local.feed, pcs_workbook=shared,
            )
            self.assertEqual(tracker_path(config), shared)
            self.assertEqual(ensure_pcs_tracker(config, **_rows()), shared)
            self.assertEqual(original.read_bytes(), original_bytes)
            self.assertEqual(shared.read_bytes(), original_bytes)

            # An existing collaborative copy wins over the old local file.
            updated = load_workbook(shared)
            updated["COACHING"]["N5"] = "Completed"
            updated.save(shared)
            updated.close()
            shared_bytes = shared.read_bytes()
            self.assertEqual(ensure_pcs_tracker(config, **_rows()), shared)
            self.assertEqual(shared.read_bytes(), shared_bytes)
            self.assertEqual(original.read_bytes(), original_bytes)

    def test_dedicated_pcs_path_requires_a_synced_local_folder(self):
        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            config = SimpleNamespace(
                reports=root / "Reports",
                pcs_workbook=root / "Missing Library" / PCS_TRACKER_FILENAME,
            )
            with self.assertRaisesRegex(RuntimeError, "folder is unavailable"):
                ensure_pcs_tracker(config, **_rows())
            self.assertFalse(config.pcs_workbook.parent.exists())

    def test_pcs_setting_only_changes_the_dedicated_path(self):
        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            config_file = root / "wfmhub.toml"
            config_file.write_text(
                '[paths]\nreports = "Reports"\nfeed = "Feed"\n\n[period]\nstart = ""\n',
                encoding="utf-8",
            )
            target = root / "Synced WFM" / PCS_TRACKER_FILENAME
            target.parent.mkdir()
            write_pcs_workbook(config_file, target)
            write_pcs_workbook(config_file, target)
            content = config_file.read_text(encoding="utf-8")
            self.assertEqual(content.count('pcs_workbook = '), 1)
            self.assertIn('reports = "Reports"', content)
            self.assertIn('feed = "Feed"', content)
            self.assertEqual(_pcs_workbook_path(root, target.as_posix()), target)
            with self.assertRaisesRegex(ConfigError, "local path"):
                _pcs_workbook_path(root, "https://example.com/tracker.xlsx")

    def test_pcs_menu_accepts_a_synced_folder_and_keeps_reports_local(self):
        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            shared_folder = root / "Synced WFM"
            shared_folder.mkdir()
            config_file = root / "wfmhub.toml"
            config_file.write_text('[paths]\nreports = "Reports"\n', encoding="utf-8")
            config = SimpleNamespace(
                file=config_file, pcs_workbook=None, reports=root / "Reports",
            )
            with (
                patch("wfmhub.cli.load_config", return_value=config),
                patch("builtins.input", side_effect=("5", str(shared_folder))),
            ):
                _pcs_menu(root)
            content = config_file.read_text(encoding="utf-8")
            self.assertIn('reports = "Reports"', content)
            self.assertIn(f'{shared_folder.as_posix()}/{PCS_TRACKER_FILENAME}', content)

    def test_text_business_dates_are_parsed_for_sqlite_rows(self):
        self.assertEqual(_as_date("2026-09-08 17:00:00"), date(2026, 9, 8))

    def test_selectable_periods_have_explicit_like_for_like_comparisons(self):
        windows = {
            item.label: item
            for item in _pcs_reporting_windows(
                date(2026, 9, 10), date(2026, 9, 8),
            )
        }
        self.assertEqual(
            (windows["Latest day"].prior_start, windows["Latest day"].prior_end),
            (date(2026, 9, 8), date(2026, 9, 8)),
        )
        self.assertEqual(
            (windows["Current week"].prior_start, windows["Current week"].prior_end),
            (date(2026, 8, 31), date(2026, 9, 3)),
        )
        self.assertEqual(
            (windows["Current MTD"].prior_start, windows["Current MTD"].prior_end),
            (date(2026, 8, 1), date(2026, 8, 10)),
        )
        available = _pcs_reporting_windows(
            date(2026, 9, 10),
            date(2026, 9, 8),
            available_start=date(2026, 7, 15),
            available_months=(date(2026, 7, 15), date(2026, 9, 10)),
        )
        labels = {item.label for item in available}
        self.assertIn("Month 2026-07", labels)
        self.assertIn("Month 2026-09", labels)
        self.assertNotIn("Month 2026-08", labels)

    def test_pcs_feeds_cover_all_available_dates_and_calendar_months(self):
        raw = sqlite3.connect(":memory:")
        conn = DatabaseConnection(raw)
        try:
            conn.execute(
                """CREATE TABLE mart.agent_pcs_day (
                       business_date DATE NOT NULL, agent_id TEXT NOT NULL,
                       agent_name TEXT, team_leader TEXT, lob TEXT, language TEXT,
                       pcs_score_sum REAL, survey_responses INTEGER,
                       pcs_participation_responses INTEGER,
                       pcs_status_calls INTEGER, low_score_responses INTEGER,
                       top_box_responses INTEGER, inbound_calls INTEGER
                   )"""
            )
            conn.executemany(
                """INSERT INTO mart.agent_pcs_day VALUES (
                       ?, '001', 'Agent One', 'TL 1', 'RSA NL', 'NL',
                       ?, 1, 1, 2, ?, ?, 3
                   )""",
                (
                    (date(2026, 7, 15), 2.0, 1, 0),
                    (date(2026, 8, 20), 3.0, 1, 0),
                    (date(2026, 9, 9), 4.0, 0, 1),
                    (date(2026, 9, 10), 5.0, 0, 1),
                ),
            )
            refreshed = datetime(2026, 9, 10, 18, 0)
            filter_rows, _lob, _agent, daily = pcs_dashboard_cache_rows(
                conn, date(2026, 9, 10), refreshed,
            )
            periods = [row[0] for row in filter_rows if row[0]]
            self.assertIn("All available", periods)
            self.assertIn("Month 2026-07", periods)
            self.assertIn("Month 2026-08", periods)
            self.assertIn("Month 2026-09", periods)

            # The all-history cards use the full period, while its chart cache
            # remains bounded to the permanent tracker's 31 plotted points.
            all_history_daily = [
                row for row in daily
                if str(row[0]).startswith("DAILY|All available|All|All|All|")
            ]
            self.assertEqual(len(all_history_daily), 31)
            self.assertEqual(all_history_daily[0][2], date(2026, 8, 11))
            self.assertEqual(all_history_daily[-1][2], date(2026, 9, 10))

            results = pcs_result_rows(
                conn, date(2026, 9, 10), 1, refreshed,
            )
            daily_dates = {
                row[1] for row in results
                if row[0] == "Daily detail" and row[3] == "AGENT DAY"
            }
            self.assertEqual(daily_dates, {
                date(2026, 7, 15), date(2026, 8, 20),
                date(2026, 9, 9), date(2026, 9, 10),
            })
            all_available = next(
                row for row in results
                if row[0] == "All available"
                and row[3] == "AGENT" and row[7] == "001"
            )
            self.assertEqual((all_available[1], all_available[2]), (
                date(2026, 7, 15), date(2026, 9, 10),
            ))
            self.assertEqual(all_available[10], 3.5)
            self.assertEqual(all_available[12], 4)
        finally:
            raw.close()

    def test_live_tracker_is_lightweight_and_created_once(self):
        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            config = SimpleNamespace(
                reports=root / "Reports",
                feed=root / "Feed",
            )
            path = ensure_pcs_tracker(config, **_rows())
            self.assertEqual(path.name, PCS_TRACKER_FILENAME)
            self.assertEqual(_tracker_contract_version(path), PCS_TRACKER_VERSION)
            with zipfile.ZipFile(path) as archive:
                self.assertIsNone(archive.testzip())
                for name in archive.namelist():
                    if name.endswith((".xml", ".rels")):
                        ElementTree.fromstring(archive.read(name))
                self.assertNotIn(b"plotVisOnly", archive.read("xl/charts/chart1.xml"))
                self.assertNotIn(b"plotVisOnly", archive.read("xl/charts/chart2.xml"))
                self.assertNotIn("xl/metadata.xml", archive.namelist())
                self.assertNotIn("xl/connections.xml", archive.namelist())
                worksheet_xml = b"".join(
                    archive.read(name) for name in archive.namelist()
                    if name.startswith("xl/worksheets/sheet") and name.endswith(".xml")
                )
                self.assertNotIn(b"_xlfn", worksheet_xml)
                self.assertNotIn(b"FILTER(", worksheet_xml)
                self.assertNotIn(b"AGGREGATE", worksheet_xml)
                self.assertNotIn(b"SUMPRODUCT", worksheet_xml)
                self.assertNotIn(b"#REF!", worksheet_xml)

            workbook = load_workbook(path, read_only=False, data_only=False)
            try:
                self.assertEqual(workbook.sheetnames, [
                    "OVERVIEW", "PERFORMANCE", "COACHING", "SETUP", "HELP",
                    "_PCS_FILTERS", "_PCS_LOB", "_PCS_AGENT", "_PCS_DAILY",
                    "_PCS_COACH", "_PCS_CALC", "_AUDIT",
                ])
                self.assertNotIn("DATA", workbook.sheetnames)
                self.assertEqual(len(workbook["OVERVIEW"]._charts), 2)
                self.assertEqual(workbook["OVERVIEW"]["A39"].value,
                                 "AGENT PERFORMANCE · CASCADING LOB → TEAM LEADER → AGENT")
                self.assertEqual(len(workbook["OVERVIEW"].data_validations.dataValidation), 4)
                self.assertEqual(workbook["OVERVIEW"]["A3"].value, "Current MTD")
                self.assertEqual(workbook["OVERVIEW"]["H3"].value, "All")
                self.assertEqual(workbook["OVERVIEW"]["O3"].value, "All")
                self.assertEqual(workbook["OVERVIEW"]["V3"].value, "All")
                self.assertTrue(workbook["OVERVIEW"]["A6"].value.startswith("=IFERROR(INDEX("))
                self.assertIn("PCS_LOB_DATA", workbook["OVERVIEW"]["A6"].value)
                self.assertNotIn("tblPcs", workbook["OVERVIEW"]["A6"].value)
                self.assertIn("PCS_LOB_DATA", workbook["_PCS_CALC"]["B2"].value)
                self.assertIn("PCS_DAILY_DATA", workbook["_PCS_CALC"]["E2"].value)
                self.assertIn("PCS_COACH_DATA", workbook["COACHING"]["A5"].value)
                for name in (
                    "PCS_LOB_DATA", "PCS_AGENT_DATA", "PCS_DAILY_DATA",
                    "PCS_COACH_DATA", "PCS_PERIOD_LIST", "PCS_LOB_LIST",
                    "PCS_TL_ACTIVE", "PCS_AGENT_ACTIVE",
                ):
                    self.assertIn(name, workbook.defined_names)
                    self.assertNotIn("#REF!", workbook.defined_names[name].attr_text)
                for name in (
                    "PCS_PERIOD_LIST", "PCS_LOB_LIST",
                    "PCS_TL_ACTIVE", "PCS_AGENT_ACTIVE",
                ):
                    formula = workbook.defined_names[name].attr_text.upper()
                    self.assertTrue("INDEX(" in formula or "OFFSET(" in formula)
                    self.assertIn("$50004", formula)
                self.assertIn("tblPcsPerformance", workbook["PERFORMANCE"].tables)
                self.assertIn("tblCoachingQueue", workbook["COACHING"].tables)
                self.assertIn("tblCoachingActions", workbook["COACHING"].tables)
                self.assertIn("tblPcsFilters", workbook["_PCS_FILTERS"].tables)
                self.assertIn("tblPcsLob", workbook["_PCS_LOB"].tables)
                self.assertIn("tblPcsAgent", workbook["_PCS_AGENT"].tables)
                self.assertIn("tblPcsDaily", workbook["_PCS_DAILY"].tables)
                self.assertIn("tblPcsCoachingView", workbook["_PCS_COACH"].tables)
                self.assertEqual(
                    tuple(cell.value for cell in workbook["_PCS_FILTERS"][4]),
                    PCS_FILTER_HEADERS,
                )
                self.assertEqual(
                    tuple(cell.value for cell in workbook["_PCS_LOB"][4]),
                    PCS_LOB_SCORECARD_HEADERS,
                )
                self.assertEqual(
                    tuple(cell.value for cell in workbook["_PCS_AGENT"][4]),
                    PCS_AGENT_SCORECARD_HEADERS,
                )
                self.assertEqual(
                    tuple(cell.value for cell in workbook["_PCS_DAILY"][4]),
                    PCS_DAILY_SCORECARD_HEADERS,
                )
                self.assertEqual(
                    tuple(cell.value for cell in workbook["PERFORMANCE"][4]),
                    PCS_RESULTS_HEADERS,
                )
                self.assertEqual(
                    tuple(cell.value for cell in workbook["COACHING"][4][:10]),
                    PCS_COACHING_HEADERS[2:12],
                )
                self.assertEqual(
                    tuple(cell.value for cell in workbook["_PCS_COACH"][4]),
                    PCS_COACHING_HEADERS,
                )
            finally:
                workbook.close()

            cached = load_workbook(path, read_only=True, data_only=True)
            try:
                self.assertEqual(cached["OVERVIEW"]["A6"].value, 4.3)
                self.assertEqual(cached["_PCS_CALC"]["B2"].value, 4.3)
                self.assertEqual(cached["COACHING"]["A5"].value, datetime(2026, 9, 8))
            finally:
                cached.close()

            state = inspect_pcs_tracker(path, config.feed / "PCS")
            self.assertIsNone(state.problem)
            self.assertTrue(state.current_template)

            saved = path.read_bytes()
            self.assertEqual(ensure_pcs_tracker(config, **_rows()), path)
            self.assertEqual(path.read_bytes(), saved)

    def test_version_migration_preserves_keyed_actions(self):
        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            config = SimpleNamespace(
                reports=root / "Reports",
                feed=root / "Feed",
            )
            with patch("wfmhub.pcs_tracker.PCS_TRACKER_VERSION", "1.0.0"):
                path = ensure_pcs_tracker(config, **_rows())
            workbook = load_workbook(path)
            sheet = workbook["COACHING"]
            sheet["L5"] = "key-migrate"
            sheet["M5"] = "call-migrate"
            sheet["N5"] = "Completed"
            sheet["O5"] = "TL 1"
            workbook.save(path)
            workbook.close()

            self.assertEqual(ensure_pcs_tracker(config, **_rows()), path)
            self.assertEqual(_tracker_contract_version(path), PCS_TRACKER_VERSION)
            archives = list((config.reports / "Archive").rglob("PCS Live Tracker_pre_*.xlsx"))
            self.assertEqual(len(archives), 1)
            migrated = load_workbook(path, read_only=True, data_only=True)
            try:
                self.assertEqual(migrated["COACHING"]["L5"].value, "key-migrate")
                self.assertEqual(migrated["COACHING"]["M5"].value, "call-migrate")
                self.assertEqual(migrated["COACHING"]["N5"].value, "Completed")
                self.assertEqual(migrated["COACHING"]["O5"].value, "TL 1")
            finally:
                migrated.close()

    def test_query_tracker_is_not_rebuilt_on_template_version_change(self):
        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            config = SimpleNamespace(reports=root / "Reports", feed=root / "Feed")
            with patch("wfmhub.pcs_tracker.PCS_TRACKER_VERSION", "1.0.0"):
                path = ensure_pcs_tracker(config, **_rows())
            with zipfile.ZipFile(path, "a") as archive:
                archive.writestr(
                    "xl/connections.xml",
                    '<connections xmlns="http://schemas.openxmlformats.org/'
                    'spreadsheetml/2006/main" count="0"/>',
                )
            before = path.read_bytes()
            self.assertEqual(ensure_pcs_tracker(config, **_rows()), path)
            self.assertEqual(path.read_bytes(), before)
            self.assertEqual(_tracker_contract_version(path), "1.0.0")

    def test_pcs_update_publishes_csv_without_rebuilding_query_tracker(self):
        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            path = root / PCS_TRACKER_FILENAME
            path.write_bytes(b"shared workbook stays untouched")
            conn = MagicMock()
            conn.execute.return_value.fetchone.return_value = (
                date(2026, 9, 1), date(2026, 9, 8),
            )
            config = SimpleNamespace(reports=root, feed=root / "Feed")
            with (
                patch("wfmhub.pcs_tracker.publish_pcs_feeds") as publish,
                patch("wfmhub.pcs_tracker._seed_dedicated_tracker", return_value=path),
                patch("wfmhub.pcs_tracker._tracker_contract_version", return_value="1.0.0"),
                patch(
                    "wfmhub.pcs_tracker.inspect_pcs_tracker",
                    return_value=SimpleNamespace(has_connections=True, query_parts=6),
                ),
                patch("wfmhub.pcs_tracker.pcs_dashboard_cache_rows") as rebuild,
            ):
                self.assertEqual(
                    build_pcs_live_tracker(
                        conn, config, date(2026, 9, 1), date(2026, 9, 8),
                    ),
                    path,
                )
            publish.assert_called_once()
            rebuild.assert_not_called()
            self.assertEqual(path.read_bytes(), b"shared workbook stays untouched")

    def test_migration_refuses_to_replace_unreadable_coaching_workbook(self):
        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            config = SimpleNamespace(reports=root / "Reports", feed=root / "Feed")
            path = config.reports / PCS_TRACKER_FILENAME
            path.parent.mkdir(parents=True)
            path.write_bytes(b"unreadable old workbook")
            with self.assertRaisesRegex(RuntimeError, "coaching actions"):
                ensure_pcs_tracker(config, **_rows())
            self.assertEqual(path.read_bytes(), b"unreadable old workbook")

    def test_latest_report_is_the_fixed_permanent_tracker(self):
        with tempfile.TemporaryDirectory() as folder:
            reports = Path(folder)
            config = SimpleNamespace(reports=reports)
            tracker = reports / PCS_TRACKER_FILENAME
            tracker.write_bytes(b"tracker")
            self.assertEqual(latest_pcs_report(config), tracker)

    def test_latest_update_loads_only_pcs_sources_and_publishes_csv(self):
        with tempfile.TemporaryDirectory() as folder:
            home = Path(folder)
            report = home / "Reports" / PCS_TRACKER_FILENAME
            report.parent.mkdir()
            report.write_bytes(b"tracker")
            config = SimpleNamespace(reports=report.parent, feed=home / "Feed")
            with (
                patch("wfmhub.cli.refresh", return_value=0) as run,
                patch("wfmhub.cli.load_config", return_value=config),
                patch("wfmhub.cli.latest_pcs_report", return_value=report),
            ):
                _build_latest_pcs_report(home)
            run.assert_called_once_with(home, None, None, ("pcs",), "pcs", False)

    def test_fast_update_uses_current_database_without_source_ingestion(self):
        with tempfile.TemporaryDirectory() as folder:
            home = Path(folder)
            report = home / "Reports" / PCS_TRACKER_FILENAME
            config = SimpleNamespace(reports=report.parent, feed=home / "Feed")
            conn = MagicMock()
            with (
                patch("wfmhub.cli.load_config", return_value=config),
                patch("wfmhub.cli.write_session", return_value=nullcontext(conn)),
                patch(
                    "wfmhub.cli._pcs_mart_period",
                    return_value=(date(2026, 8, 1), date(2026, 9, 8)),
                ),
                patch("wfmhub.cli.build_report_pack", return_value=report) as build,
            ):
                self.assertEqual(_build_pcs_from_database(home), report)
            build.assert_called_once_with(
                "pcs", conn, config, date(2026, 8, 1), date(2026, 9, 8),
            )

    def test_install_uses_desktop_excel_helper_once(self):
        with tempfile.TemporaryDirectory() as folder:
            home = Path(folder)
            report = home / "Reports" / PCS_TRACKER_FILENAME
            report.parent.mkdir()
            report.write_bytes(b"tracker")
            config = SimpleNamespace(reports=report.parent, feed=home / "Feed")
            with (
                patch("wfmhub.cli.load_config", return_value=config),
                patch("wfmhub.cli.latest_pcs_report", return_value=report),
                patch(
                    "wfmhub.cli.inspect_pcs_tracker",
                    return_value=SimpleNamespace(
                        current_template=True, problem=None,
                        has_connections=False, query_parts=0,
                    ),
                ),
                patch("wfmhub.cli.run_pcs_excel_action", return_value="ready") as run,
            ):
                _sync_pcs_workbook(home)
            run.assert_called_once_with(config, report, "Sync")

    def test_install_repairs_an_outdated_tracker_before_excel_automation(self):
        with tempfile.TemporaryDirectory() as folder:
            home = Path(folder)
            old_report = home / "Reports" / PCS_TRACKER_FILENAME
            repaired_report = home / "Reports" / PCS_TRACKER_FILENAME
            config = SimpleNamespace(reports=old_report.parent, feed=home / "Feed")
            old_state = SimpleNamespace(
                current_template=False, problem="#REF!",
                has_connections=False, query_parts=0,
            )
            repaired_state = SimpleNamespace(
                current_template=True, problem=None,
                has_connections=False, query_parts=0,
            )
            with (
                patch("wfmhub.cli.load_config", return_value=config),
                patch("wfmhub.cli.latest_pcs_report", return_value=old_report),
                patch(
                    "wfmhub.cli.inspect_pcs_tracker",
                    side_effect=(old_state, repaired_state),
                ),
                patch(
                    "wfmhub.cli._build_pcs_from_database",
                    return_value=repaired_report,
                ) as repair,
                patch("wfmhub.cli.run_pcs_excel_action", return_value="ready") as run,
            ):
                _sync_pcs_workbook(home)
            repair.assert_called_once_with(home)
            run.assert_called_once_with(config, repaired_report, "Sync")

    def test_legacy_sync_refuses_query_workbook_without_rebuild(self):
        with tempfile.TemporaryDirectory() as folder:
            home = Path(folder)
            report = home / "Reports" / PCS_TRACKER_FILENAME
            report.parent.mkdir()
            report.write_bytes(b"tracker")
            config = SimpleNamespace(reports=report.parent, feed=home / "Feed")
            state = SimpleNamespace(
                current_template=False, problem="#REF!",
                has_connections=True, query_parts=6,
            )
            with (
                patch("wfmhub.cli.load_config", return_value=config),
                patch("wfmhub.cli.latest_pcs_report", return_value=report),
                patch("wfmhub.cli.inspect_pcs_tracker", return_value=state),
                patch("wfmhub.cli._build_pcs_from_database") as rebuild,
                patch("wfmhub.cli.run_pcs_excel_action") as sync,
            ):
                with self.assertRaisesRegex(RuntimeError, "Excel queries/connections"):
                    _sync_pcs_workbook(home)
            rebuild.assert_not_called()
            sync.assert_not_called()

    def test_pcs_refresh_uses_targeted_model_without_all_shared_feeds(self):
        home = Path("/test/wfmhub")
        config = SimpleNamespace(
            business_rules=Path("rules.toml"),
            queue_mapping=Path("queues.csv"),
        )
        conn = MagicMock()
        ingest = SimpleNamespace(
            loaded=0, skipped=2, failed=0, scoped_out=0, errors=[],
        )
        model = ModelSummary(
            start=date(2026, 8, 1), end=date(2026, 9, 8), pcs_rows=12,
        )
        rulebook = SimpleNamespace(version="rules", sha256="a" * 64)
        mapping = SimpleNamespace(sha256="b" * 64)
        report = Path(f"/test/wfmhub/Reports/{PCS_TRACKER_FILENAME}")
        with (
            patch("wfmhub.cli.load_config", return_value=config),
            patch("wfmhub.cli._logging"),
            patch("wfmhub.cli.ProgressBar", return_value=MagicMock()),
            patch("wfmhub.cli.write_session", return_value=nullcontext(conn)),
            patch("wfmhub.cli.ingest_all", return_value=ingest),
            patch("wfmhub.cli.refresh_pcs_models", return_value=model) as pcs_models,
            patch("wfmhub.cli.refresh_models") as all_models,
            patch("wfmhub.cli.publish_shared_feeds") as all_feeds,
            patch("wfmhub.cli.build_report_pack", return_value=report) as build,
            patch("wfmhub.cli.load_rulebook", return_value=rulebook),
            patch("wfmhub.cli.load_queue_mapping", return_value=mapping),
        ):
            result = refresh(home, None, None, ("pcs",), "pcs", False)
        self.assertEqual(result, 0)
        pcs_models.assert_called_once()
        all_models.assert_not_called()
        all_feeds.assert_not_called()
        build.assert_called_once_with(
            "pcs", conn, config, model.start, model.end, service_profile=None,
        )


if __name__ == "__main__":
    unittest.main()
