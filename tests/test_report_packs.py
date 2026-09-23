from __future__ import annotations

import tempfile
import unittest
from datetime import datetime
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import patch

from openpyxl import load_workbook

from wfmhub.design import COLORS, REPORT_DESIGN_ID, REPORT_DESIGN_VERSION
from wfmhub.report_packs import (
    REPORT_PACK_KEYS,
    REPORT_PACKS,
    ReportPublishError,
    publish_report,
    report_current_path,
)
from wfmhub.reports import COLORS as REPORT_COLORS


REPO = Path(__file__).resolve().parents[1]


class ReportPackTests(unittest.TestCase):
    def test_report_design_has_one_palette_and_reference_workbook(self):
        self.assertIs(REPORT_COLORS, COLORS)
        self.assertEqual((REPORT_DESIGN_ID, REPORT_DESIGN_VERSION), (
            "WFMHUB-DESIGN", "1.0.0",
        ))
        reference = REPO / "docs" / "WFMHub Report Design Reference.xlsx"
        self.assertTrue(reference.is_file())
        workbook = load_workbook(reference, read_only=False, data_only=False)
        try:
            self.assertEqual(workbook["DESIGN SYSTEM"]["A1"].value, "WFMHUB REPORT SYSTEM")
            self.assertGreaterEqual(len(workbook.sheetnames), 8)
        finally:
            workbook.close()

    def test_registry_contains_only_current_products(self):
        self.assertEqual(REPORT_PACK_KEYS, tuple(REPORT_PACKS))
        self.assertEqual(set(REPORT_PACKS), {
            "absence", "corrections", "pcs", "bonus", "service",
            "realisations", "staffing",
        })
        self.assertEqual(REPORT_PACKS["pcs"].current_filename, "PCS Live Tracker.xlsx")
        self.assertEqual(
            REPORT_PACKS["absence"].current_filename,
            "Final Absenteeism & Shrinkage.xlsx",
        )

    def test_current_reports_are_flat_and_previous_copy_is_archived(self):
        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            config = SimpleNamespace(reports=root / "Reports")
            current = report_current_path(config, "service")
            current.parent.mkdir(parents=True)
            current.write_bytes(b"old")
            partial = current.with_name("RTM Daily Control.partial.xlsx")
            partial.write_bytes(b"new")
            publish_report(config, "service", partial, current, datetime(2026, 9, 4, 10, 30))
            self.assertEqual(current.read_bytes(), b"new")
            archives = list((root / "Reports" / "Archive" / "2026-09-04").glob("*.xlsx"))
            self.assertEqual(len(archives), 1)
            self.assertEqual(archives[0].read_bytes(), b"old")

    def test_locked_workbook_has_an_actionable_publish_error(self):
        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            partial = root / "PCS Live Tracker.partial.xlsx"
            target = root / "PCS Live Tracker.xlsx"
            partial.write_bytes(b"complete replacement")
            with patch.object(Path, "replace", side_effect=PermissionError("access denied")):
                with self.assertRaisesRegex(
                    ReportPublishError, "open in Excel.*locked by OneDrive",
                ):
                    publish_report(SimpleNamespace(), "custom", partial, target, datetime.now())


if __name__ == "__main__":
    unittest.main()
