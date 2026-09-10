import csv
import json
import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook

from tools.check_power_platform_absence_contract import inspect_fte_contract
from wfmhub.ingestion import parse_fte


ROOT = Path(__file__).resolve().parents[1]
KIT = ROOT / "power-platform" / "wfm-absence-app"
TEMPLATE = ROOT / "templates" / "FTE Count.xlsx"


class PowerPlatformAbsenceContractTests(unittest.TestCase):
    def test_standard_fte_template_matches_app_contract(self):
        result = inspect_fte_contract(TEMPLATE)
        self.assertTrue(result["ok"], result["errors"])
        self.assertEqual(
            set(result["tables"]),
            {"tblFTEAgents", "tblFTEPTO", "tblFTEAway"},
        )

    def test_english_is_the_default_language(self):
        schema = json.loads((KIT / "contracts" / "sharepoint-lists.json").read_text())
        self.assertEqual(schema["defaultLanguage"], "en")
        default_variable = next(
            item
            for item in schema["environmentVariables"]
            if item["schemaName"] == "wfm_DefaultLanguage"
        )
        self.assertEqual(default_variable["default"], "en")
        prototype = (ROOT / "docs" / "prototypes" / "wfm_absence_app_form.html").read_text()
        self.assertIn('<html lang="en">', prototype)
        self.assertIn('<span class="active">EN</span>', prototype)

    def test_every_interface_label_has_english_and_french(self):
        with (KIT / "locales" / "labels.csv").open(newline="", encoding="utf-8") as handle:
            rows = list(csv.DictReader(handle))
        self.assertGreater(len(rows), 40)
        self.assertEqual(len({row["key"] for row in rows}), len(rows))
        self.assertTrue(all(row["en"].strip() for row in rows))
        self.assertTrue(all(row["fr"].strip() for row in rows))

    def test_office_script_never_targets_agent_table(self):
        script = (KIT / "office-scripts" / "WFM_Write_Absence.ts").read_text()
        self.assertNotIn('getTable(workbook, "tblFTEAgents")', script)
        self.assertIn('"tblFTEPTO"', script)
        self.assertIn('"tblFTEAway"', script)

    def test_writer_shaped_rows_are_accepted_by_wfmhub(self):
        with tempfile.TemporaryDirectory() as folder:
            target = Path(folder) / "FTE Count.xlsx"
            workbook = load_workbook(TEMPLATE)
            agent = workbook["Agent"]
            agent_values = [
                "00123", "Active", "Jane Agent", "TL One", "Ops One",
                "RSA NL", "NL", "NL", "Site", "City", 1, None,
            ]
            for column, value in enumerate(agent_values, 1):
                agent.cell(2, column).value = value

            pto = workbook["PTO"]
            pto_values = [
                "00123", "Jane Agent", "2026-09-15", "2026-09-16",
                "Full day", "", "", "PTO", "Approved", "Approved by WFM",
            ]
            for column, value in enumerate(pto_values, 1):
                pto.cell(2, column).value = value

            away = workbook["Away"]
            away_values = [
                "00123", "Jane Agent", "2026-09-20", "", "Long sickness",
                "Active", "Open case",
            ]
            for column, value in enumerate(away_values, 1):
                away.cell(2, column).value = value
            workbook.save(target)
            workbook.close()

            parsed = parse_fte(target, "power-platform-test")
            self.assertEqual(parsed.rejected, [])
            rows = parsed.tables["raw.fte_time_off"]
            self.assertEqual(len(rows), 2)
            self.assertEqual(rows[0]["agent_id"], "00123")
            self.assertEqual(rows[0]["record_status"], "APPROVED")
            self.assertEqual(rows[0]["absence_type"], "PTO")
            self.assertEqual(rows[1]["source_kind"], "AWAY")
            self.assertEqual(rows[1]["record_status"], "ACTIVE")
            self.assertIsNone(rows[1]["end_date"])


if __name__ == "__main__":
    unittest.main()
