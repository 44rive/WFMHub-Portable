from __future__ import annotations

import sqlite3
import tempfile
import unittest
from datetime import date, datetime
from pathlib import Path
from types import SimpleNamespace

from openpyxl import load_workbook

from wfmhub.database import DatabaseConnection
from wfmhub.decision_products import _break_meal_control_rows
from wfmhub.reports import ExcelReport
from wfmhub.shift_view import add_review_board


class AttendanceControlTests(unittest.TestCase):
    def test_break_and_meal_overruns_are_evidence_gated(self):
        conn = DatabaseConnection(sqlite3.connect(":memory:"))
        conn.execute(
            """CREATE TABLE mart.attendance_agent_day (
                   business_date DATE, lob VARCHAR, team_leader VARCHAR,
                   agent_id VARCHAR, agent_name VARCHAR,
                   scheduled_start TIMESTAMP, scheduled_end TIMESTAMP,
                   planned_work_minutes INTEGER, status_covered_minutes INTEGER,
                   actual_evidence VARCHAR, assignment_type VARCHAR
               )"""
        )
        conn.execute(
            """CREATE TABLE mart.shift_timeline_segment (
                   business_date DATE, agent_id VARCHAR, actual_category VARCHAR,
                   segment_start TIMESTAMP, segment_end TIMESTAMP,
                   mismatch_type VARCHAR
               )"""
        )
        day = date(2026, 8, 1)
        conn.executemany(
            "INSERT INTO mart.attendance_agent_day VALUES (?,?,?,?,?,?,?,?,?,?,?)",
            [
                (day, "RSA NL", "TL", "100", "Agent 100", datetime(2026, 8, 1, 8), datetime(2026, 8, 1, 16), 480, 480, "AGENT_STATUS", "Working"),
                (day, "RSA NL", "TL", "200", "Agent 200", datetime(2026, 8, 1, 8), datetime(2026, 8, 1, 16), 480, 120, "AGENT_STATUS", "Working"),
            ],
        )
        conn.executemany(
            "INSERT INTO mart.shift_timeline_segment VALUES (?,?,?,?,?,?)",
            [
                (day, "100", "Break", datetime(2026, 8, 1, 9), datetime(2026, 8, 1, 9, 20), "MATCH"),
                (day, "100", "Break", datetime(2026, 8, 1, 11), datetime(2026, 8, 1, 11, 20), "MATCH"),
                (day, "100", "Lunch", datetime(2026, 8, 1, 12), datetime(2026, 8, 1, 12, 50), "MATCH"),
            ],
        )
        config = SimpleNamespace(rules=SimpleNamespace(
            minimum_status_coverage=0.80, break_minutes=30, lunch_minutes=45,
        ))

        headers, rows = _break_meal_control_rows(conn, config, day, day)
        by_agent = {row[headers.index("agent_id")]: row for row in rows}
        first = by_agent["100"]
        self.assertEqual(first[headers.index("break_minutes")], 40)
        self.assertEqual(first[headers.index("break_overrun_minutes")], 10)
        self.assertEqual(first[headers.index("meal_minutes")], 50)
        self.assertEqual(first[headers.index("meal_overrun_minutes")], 5)
        self.assertEqual(first[headers.index("alert")], "BREAK & MEAL EXCEEDED")
        self.assertEqual(
            by_agent["200"][headers.index("alert")], "INSUFFICIENT EVIDENCE",
        )
        conn.close()

    def test_review_board_distinguishes_owned_other_and_tolerance_gaps(self):
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / "review.xlsx"
            report = ExcelReport(path)
            day = date(2026, 8, 1)
            headers = [
                "gap_id", "business_date", "agent_id", "agent_name",
                "team_leader", "lob", "detected_issue", "exact_start",
                "exact_end", "minutes", "suggested_activity",
                "decision_category", "decision_status", "reviewed_by",
                "comment", "reviewed_date",
            ]
            rows = [
                ("g1", day, "100", "Agent", "TL", "RSA NL", "Gap", datetime(2026, 8, 1, 8, 10), datetime(2026, 8, 1, 8, 20), 10, "Review", None, "Open", None, None, None),
                ("g2", day, "100", "Agent", "TL", "RSA NL", "Gap", datetime(2026, 8, 1, 9), datetime(2026, 8, 1, 9, 10), 10, "Review", None, "Open", None, None, None),
            ]
            base = {
                "business_date": day, "agent_id": "100", "agent_name": "Agent",
                "team_leader": "TL", "lob": "RSA NL", "language": "NL",
                "scheduled_start": datetime(2026, 8, 1, 8),
                "scheduled_end": datetime(2026, 8, 1, 10),
                "planned_state": "Working", "observed_source": "AGENT_STATUS",
            }
            segments = [
                {**base, "segment_start": datetime(2026, 8, 1, 8, 10), "segment_end": datetime(2026, 8, 1, 8, 20), "actual_category": "Logged Off", "mismatch_type": "LOGGED_OFF", "is_gap": True},
                {**base, "segment_start": datetime(2026, 8, 1, 8, 40), "segment_end": datetime(2026, 8, 1, 8, 44), "actual_category": "Logged Off", "mismatch_type": "LOGGED_OFF", "is_gap": True},
                {**base, "segment_start": datetime(2026, 8, 1, 9), "segment_end": datetime(2026, 8, 1, 9, 10), "actual_category": "Logged Off", "mismatch_type": "LOGGED_OFF", "is_gap": True},
            ]
            add_review_board(report, headers, rows, segments, day, day, ["Late"])
            report.close()

            workbook = load_workbook(path, data_only=True)
            try:
                sheet = workbook["REVIEW BOARD"]
                columns = {cell.value: cell.column for cell in sheet[4]}
                first_row = 5
                self.assertEqual(sheet.cell(first_row, columns["08:00"]).value, "Gap")
                self.assertTrue(sheet.cell(first_row, columns["08:00"]).fill.fgColor.rgb.endswith("B42318"))
                self.assertEqual(sheet.cell(first_row, columns["08:30"]).value, "Tolerance")
                self.assertTrue(sheet.cell(first_row, columns["08:30"]).fill.fgColor.rgb.endswith("EEF1F4"))
                self.assertEqual(sheet.cell(first_row, columns["09:00"]).value, "Gap")
                self.assertTrue(sheet.cell(first_row, columns["09:00"]).fill.fgColor.rgb.endswith("FDE7E5"))
            finally:
                workbook.close()


if __name__ == "__main__":
    unittest.main()
