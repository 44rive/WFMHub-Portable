"""Read and calculate the governed Bonus Matrix source contract."""

from __future__ import annotations

from collections import Counter
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from typing import Any, Iterable

from openpyxl import load_workbook


@dataclass(frozen=True)
class BonusRule:
    population: str
    kpi: str
    direction: str
    tier1_bonus: float
    tier1_target: float
    tier2_bonus: float
    tier2_target: float


@dataclass
class BonusResult:
    core_ready: bool = False
    eligibility: str = ""
    earned: tuple[float, ...] = ()
    gross: float | None = None
    malus: float | None = None
    final: float | None = None
    reference: float | None = None
    proration: float | None = None
    scenario: float | None = None
    release: float | None = None
    status: str = ""
    issue: str = ""


DEFAULT_POLICIES = [
    ("Malus method", "Proportional", "Proportional | Percentage points", "HR / Compensation"),
    ("Extra PCS treatment", "Additive", "Additive | Replacement", "Operations / HR"),
    ("Target bonus rate basis", "Monthly", "Monthly | Annual", "HR / Payroll"),
    ("Proration denominator", "Calendar days", "Calendar days | Working days | Planned days", "HR / Payroll"),
    ("Absence treatment", "KPI only", "KPI only | Eligibility only | Both", "HR / Operations"),
    ("Achievement cap", 1.30, "0% to 200%", "HR / Compensation"),
    ("Rounding", "Centime", "Centime | Whole MAD", "Payroll"),
    ("Employee eligibility", "Active eligible population", "Policy reference required", "HR"),
]


def _clean(value: Any) -> str:
    return " ".join(str(value or "").strip().split())


def _number(value: Any) -> float | None:
    if value in (None, ""):
        return None
    try:
        return float(value)
    except (TypeError, ValueError):
        return None


def _header_row(ws, required: str, maximum: int = 20) -> int:
    for row in range(1, min(ws.max_row, maximum) + 1):
        if any(_clean(ws.cell(row, col).value).casefold() == required.casefold()
               for col in range(1, ws.max_column + 1)):
            return row
    raise ValueError(f"Could not find '{required}' header in sheet {ws.title}")


def _cached_result_total(source: Path) -> tuple[float | None, int]:
    """Read the source workbook's cached Results payout for reconciliation."""

    workbook = load_workbook(source, read_only=True, data_only=True, keep_links=False)
    try:
        if "Results" not in workbook.sheetnames:
            return None, 0
        sheet = workbook["Results"]
        header_row = _header_row(sheet, "Agent ID")
        headers = [
            _clean(sheet.cell(header_row, column).value)
            for column in range(1, sheet.max_column + 1)
        ]
        if "Final Payout" not in headers:
            return None, 0
        payout_index = headers.index("Final Payout")
        agent_index = headers.index("Agent ID")
        values: list[float] = []
        for row in sheet.iter_rows(min_row=header_row + 1, values_only=True):
            if agent_index >= len(row) or not _clean(row[agent_index]):
                continue
            payout = _number(row[payout_index] if payout_index < len(row) else None)
            if payout is not None:
                values.append(payout)
        return sum(values), len(values)
    finally:
        workbook.close()


def _read_bonus_source(path: Path) -> tuple[list[dict[str, Any]], list[BonusRule], list[tuple[Any, ...]]]:
    if not path.exists():
        raise FileNotFoundError(path)
    workbook = load_workbook(path, read_only=True, data_only=True, keep_links=False)
    try:
        if "Raw_Data" not in workbook.sheetnames or "KPI_Config" not in workbook.sheetnames:
            raise ValueError("Bonus source must contain Raw_Data and KPI_Config sheets")
        raw = workbook["Raw_Data"]
        row = _header_row(raw, "Agent ID")
        headers = [_clean(raw.cell(row, col).value) for col in range(1, raw.max_column + 1)]
        rows: list[dict[str, Any]] = []
        for values in raw.iter_rows(min_row=row + 1, values_only=True):
            record = {headers[index]: value for index, value in enumerate(values) if index < len(headers)}
            if not _clean(record.get("Agent ID")):
                continue
            rows.append(record)

        config = workbook["KPI_Config"]
        config_row = _header_row(config, "Population")
        rules: list[BonusRule] = []
        for values in config.iter_rows(min_row=config_row + 1, values_only=True):
            if not _clean(values[0] if values else None) or not _clean(values[1] if len(values) > 1 else None):
                continue
            rules.append(BonusRule(
                _clean(values[0]), _clean(values[1]), _clean(values[2]).upper(),
                float(values[3] or 0), float(values[4] or 0),
                float(values[5] or 0), float(values[6] or 0),
            ))

        policies: list[tuple[Any, ...]] = []
        if "Policy_Decisions" in workbook.sheetnames:
            policy = workbook["Policy_Decisions"]
            policy_row = _header_row(policy, "Policy")
            for values in policy.iter_rows(min_row=policy_row + 1, values_only=True):
                if not _clean(values[0] if values else None):
                    continue
                policies.append(tuple(values[:7]))
        if not policies:
            policies = [(name, value, allowed, "", owner, "To validate", "")
                        for name, value, allowed, owner in DEFAULT_POLICIES]
        return rows, rules, policies
    finally:
        workbook.close()


def _bonus_value(record: dict[str, Any], *names: str) -> Any:
    for name in names:
        if name in record:
            return record[name]
    return None


def _bonus_period(records: list[dict[str, Any]]) -> str:
    periods = Counter(_clean(_bonus_value(row, "Period")) for row in records)
    periods.pop("", None)
    return periods.most_common(1)[0][0] if periods else datetime.now().strftime("%Y-%m")


def _policy_map(policies: Iterable[tuple[Any, ...]]) -> dict[str, Any]:
    return {_clean(row[0]): row[1] for row in policies}


def _earned(actual: float, rule: BonusRule) -> float:
    first = actual <= rule.tier1_target if rule.direction == "L" else actual >= rule.tier1_target
    second = actual <= rule.tier2_target if rule.direction == "L" else actual >= rule.tier2_target
    return rule.tier1_bonus if first else rule.tier2_bonus if second else 0.0


def _calculate_bonus(
    record: dict[str, Any],
    rules: list[BonusRule],
    policies: dict[str, Any],
    policies_ready: bool,
) -> BonusResult:
    population = _clean(_bonus_value(record, "Population"))
    actuals = [
        _number(_bonus_value(record, "AHT")),
        _number(_bonus_value(record, "Productivity")),
        _number(_bonus_value(record, "PCS Score")),
        _number(_bonus_value(record, "PCS % Participation", "PCS % (Participation)")),
        _number(_bonus_value(record, "QM")),
        _number(_bonus_value(record, "Abs%")),
    ]
    reference_override = _number(_bonus_value(record, "Reference Bonus Override")) or 0
    salary = _number(_bonus_value(record, "Monthly Fixed Salary")) or 0
    rate = _number(_bonus_value(record, "Target Bonus Rate")) or 0
    eligible_days = _number(_bonus_value(record, "Eligible Days"))
    scheduled_days = _number(_bonus_value(record, "Scheduled Days"))
    rule_map = {(rule.population, rule.kpi): rule for rule in rules}
    required = ("AHT", "Productivity", "PCS Score", "PCS % (Participation)", "QM", "Abs%")
    result = BonusResult()
    result.core_ready = (
        all(value is not None for value in actuals)
        and all((population, name) in rule_map for name in required)
        and (reference_override > 0 or (salary > 0 and rate > 0))
        and eligible_days is not None and scheduled_days is not None
        and 0 <= eligible_days <= scheduled_days and scheduled_days > 0
        and all(0 <= actuals[index] <= 1 for index in (3, 4, 5))
    )
    if not result.core_ready:
        result.status = "BLOCKED - INPUT"
        result.issue = "Missing, invalid, or unconfigured required input"
        return result
    scores = [_earned(float(value), rule_map[(population, name)]) for value, name in zip(actuals, required)]
    extra_rule = rule_map.get((population, "Extra Bonus (PCS Score)"))
    extra = _earned(float(actuals[2]), extra_rule) if extra_rule else 0.0
    absence_treatment = _clean(policies.get("Absence treatment", "KPI only"))
    if absence_treatment == "Eligibility only":
        scores[5] = 0.0
    absence_target = rule_map[(population, "Abs%")].tier1_target
    employment = _clean(_bonus_value(record, "Employment Status")) or "Not supplied"
    absence_eligible = float(actuals[5]) <= absence_target
    if employment.casefold() != "active":
        result.eligibility = "REVIEW"
    elif absence_treatment in {"Eligibility only", "Both"} and not absence_eligible:
        result.eligibility = "INELIGIBLE ABS"
    else:
        result.eligibility = "ELIGIBLE"
    extra_method = _clean(policies.get("Extra PCS treatment", "Additive"))
    base = sum(scores)
    gross = base + extra if extra_method == "Additive" else base - scores[2] + max(scores[2], extra)
    cap = _number(policies.get("Achievement cap")) or 1.30
    gross = min(gross, cap)
    voc = int(_number(_bonus_value(record, "VOC Detractor Count")) or 0)
    malus = {0: 0.0, 1: 0.10, 2: 0.20, 3: 0.50, 4: 0.75}.get(voc, 1.0)
    final = max(0.0, gross - malus) if _clean(policies.get("Malus method")) == "Percentage points" else gross * (1 - malus)
    reference = reference_override or salary * rate / (12 if _clean(policies.get("Target bonus rate basis")) == "Annual" else 1)
    proration = min(1.0, float(eligible_days) / float(scheduled_days))
    digits = 0 if _clean(policies.get("Rounding")) == "Whole MAD" else 2
    scenario = round(reference * proration * final, digits)
    data_status = _clean(_bonus_value(record, "Data Status"))
    result.earned = (*scores, extra)
    result.gross, result.malus, result.final = gross, malus, final
    result.reference, result.proration, result.scenario = reference, proration, scenario
    result.release = scenario if policies_ready and data_status == "VALIDATED" and result.eligibility == "ELIGIBLE" else None
    if result.eligibility != "ELIGIBLE":
        result.status = "BLOCKED - ELIGIBILITY"
    elif data_status != "VALIDATED":
        result.status = "BLOCKED - DATA"
    elif not policies_ready:
        result.status = "BLOCKED - POLICY"
    else:
        result.status = "READY"
    result.issue = "" if result.status == "READY" else result.status
    return result
