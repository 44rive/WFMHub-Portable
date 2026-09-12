"""Permanent, lightweight PCS Power Query tracker lifecycle.

Python and SQLite calculate the governed PCS products and publish fixed CSV
feeds. Desktop Excel uses Power Query only as transport. The permanent workbook
contains lightweight report tables and a permanent human-owned coaching action
log; raw call-leg data is never loaded to a worksheet.
"""

from __future__ import annotations

import os
import shutil
from datetime import date, datetime
from pathlib import Path
from typing import Any, Sequence
from xml.etree import ElementTree
from zipfile import BadZipFile, ZipFile

from openpyxl import load_workbook

from .config import Config
from .database import DatabaseConnection
from .design import COLORS, REPORT_DESIGN_ID, REPORT_DESIGN_VERSION
from .excel_layout import (
    ACTION_FIRST_ROW,
    ACTION_HEADER_ROW,
    ACTION_SECTION_ROW,
    configure_v2_cell_canvas,
    make_v2_formats,
    style_v2_chart,
    write_v2_header,
    write_v2_kpis,
)
from .metrics import load_metric_catalog
from .reports import ExcelReport
from .shared_feeds import (
    PCS_AGENT_SCORECARD_HEADERS,
    PCS_COACHING_HEADERS,
    PCS_DAILY_SCORECARD_HEADERS,
    PCS_FILTER_HEADERS,
    PCS_LOB_SCORECARD_HEADERS,
    PCS_RESULTS_HEADERS,
    pcs_coaching_cache_rows,
    pcs_dashboard_cache_rows,
    pcs_result_rows,
    publish_pcs_feeds,
)


PCS_TRACKER_FILENAME = "PCS Live Tracker.xlsx"
PCS_TRACKER_VERSION = "2026.11.4"
COACHING_ACTION_HEADERS = (
    "Coaching Key", "Call ID", "Coaching Status", "Coach",
    "Coaching Date", "Due Date", "Coaching Comment",
)


def tracker_path(config: Config) -> Path:
    return (config.reports / PCS_TRACKER_FILENAME).resolve()


def latest_pcs_report(config: Config) -> Path | None:
    path = tracker_path(config)
    return path if path.is_file() else None


def open_workbook(path: Path) -> None:
    if not path.is_file():
        raise FileNotFoundError(path)
    if os.name != "nt":
        raise RuntimeError(f"Open this workbook manually: {path}")
    os.startfile(path)  # type: ignore[attr-defined]


def _as_date(value: Any) -> date | None:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if value in (None, ""):
        return None
    try:
        return date.fromisoformat(str(value)[:10])
    except ValueError:
        return None


def _formula_cache(value: Any) -> Any:
    """Return a scalar Excel can safely store as a formula cached value."""

    if value is None:
        return ""
    if isinstance(value, datetime):
        days = (value.date() - date(1899, 12, 30)).days
        seconds = (
            value.hour * 3600 + value.minute * 60 + value.second
            + value.microsecond / 1_000_000
        )
        return days + seconds / 86_400
    if isinstance(value, date):
        return (value - date(1899, 12, 30)).days
    return value


def _read_actions_from(path: Path) -> list[dict[str, Any]]:
    """Read keyed action rows from every supported PCS tracker layout."""

    if not path.is_file():
        return []
    try:
        workbook = load_workbook(
            path, read_only=True, data_only=True, keep_links=False,
        )
    except Exception:
        return []
    try:
        if "COACHING" not in workbook.sheetnames:
            return []
        sheet = workbook["COACHING"]
        header_row = None
        headers: dict[str, int] = {}
        for row_number in range(1, min(40, sheet.max_row) + 1):
            candidate = {
                str(cell.value or "").strip(): cell.column
                for cell in sheet[row_number]
                if cell.value not in (None, "")
            }
            if {"Coaching Key", "Coaching Status", "Coach"} <= set(candidate):
                header_row = row_number
                headers = candidate
                break
        if header_row is None:
            return []
        output: list[dict[str, Any]] = []
        seen: set[str] = set()
        for values in sheet.iter_rows(min_row=header_row + 1, values_only=True):
            key_column = headers["Coaching Key"]
            key = values[key_column - 1] if key_column <= len(values) else None
            key_text = str(key or "").strip()
            if not key_text or key_text in seen:
                continue
            output.append({
                header: values[column - 1] if column <= len(values) else None
                for header, column in headers.items()
                if header in COACHING_ACTION_HEADERS
            })
            seen.add(key_text)
        return output
    finally:
        workbook.close()


def _tracker_contract_version(path: Path) -> str | None:
    """Read the tracker contract without asking Excel to repair the file."""

    try:
        with ZipFile(path) as archive:
            root = ElementTree.fromstring(archive.read("docProps/custom.xml"))
    except (BadZipFile, KeyError, OSError, ElementTree.ParseError):
        return None
    for prop in root:
        if prop.attrib.get("name") != "WFMHub PCS Live Tracker":
            continue
        for child in prop:
            return child.text
    return None


def _legacy_actions(config: Config) -> list[tuple[Any, ...]]:
    """Collect unique keyed actions during the one-time architecture migration."""

    candidates = (
        tracker_path(config),
        config.reports / "PCS Coaching Log.xlsx",
        config.reports / "PCS Operational Tracker.xlsx",
    )
    selected: dict[str, dict[str, Any]] = {}
    for path in candidates:
        for row in _read_actions_from(path):
            key = str(row.get("Coaching Key") or "").strip()
            if key and key not in selected:
                selected[key] = row
    return [
        tuple(row.get(header) for header in COACHING_ACTION_HEADERS)
        for row in selected.values()
    ]


def _starter_rows(
    rows: Sequence[Sequence[Any]],
    headers: Sequence[str],
) -> list[tuple[Any, ...]]:
    return [tuple(row) for row in rows] or [tuple(None for _ in headers)]


def _cell_format(report: ExcelReport, header: str, value: Any, *, editable: bool = False):
    if editable:
        return report.editable_date if "Date" in header else report.editable
    if header in {"Date", "As Of Date", "Period Start", "Period End", "Data Through"}:
        return report.date
    if header in {"Call Start", "Feed Refreshed At"}:
        return report.datetime
    if "Participation" in header or header.endswith("Rate"):
        return report.percent
    if any(token in header for token in ("PCS", "Change")) and not any(
        token in header for token in ("Status", "Rule")
    ):
        return report.decimal
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        return report.integer
    return report.body


def _write_table(
    report: ExcelReport,
    worksheet,
    *,
    first_row: int,
    first_column: int,
    table_name: str,
    headers: Sequence[str],
    rows: Sequence[Sequence[Any]],
    editable_headers: set[str] | None = None,
) -> None:
    values = _starter_rows(rows, headers)
    editable_headers = editable_headers or set()
    for row_offset, row in enumerate(values, 1):
        for column_offset, value in enumerate(row):
            header = headers[column_offset]
            worksheet.write(
                first_row + row_offset,
                first_column + column_offset,
                value,
                _cell_format(
                    report,
                    header,
                    value,
                    editable=header in editable_headers,
                ),
            )
    worksheet.add_table(
        first_row,
        first_column,
        first_row + len(values),
        first_column + len(headers) - 1,
        {
            "name": table_name,
            "style": "Table Style Light 9",
            "columns": [
                {"header": header, "header_format": report.header}
                for header in headers
            ],
        },
    )


def _query_sheet(
    report: ExcelReport,
    *,
    name: str,
    title: str,
    subtitle: str,
    table_name: str,
    headers: Sequence[str],
    rows: Sequence[Sequence[Any]],
    hidden: bool = False,
):
    worksheet = report.workbook.add_worksheet(name)
    worksheet.hide_gridlines(2)
    worksheet.set_tab_color(COLORS["teal"])
    worksheet.merge_range(0, 0, 0, len(headers) - 1, title, report.title)
    worksheet.merge_range(1, 0, 1, len(headers) - 1, subtitle, report.subtitle)
    worksheet.set_row(0, 30)
    worksheet.set_row(1, 19)
    worksheet.freeze_panes(4, 0)
    _write_table(
        report,
        worksheet,
        first_row=3,
        first_column=0,
        table_name=table_name,
        headers=headers,
        rows=rows,
    )
    for column, header in enumerate(headers):
        width = 15
        if header in {"Agent", "Agent Selector", "Team Leader"}:
            width = 24
        elif header in {"Customer Comment", "Coaching Key"}:
            width = 40
        elif header in {"Call ID"}:
            width = 25
        elif "Date" in header or header.endswith(" At"):
            width = 19
        worksheet.set_column(column, column, width)
    worksheet.set_zoom(85)
    if hidden:
        worksheet.hide()
    return worksheet


def _cached(row: Sequence[Any] | None, header: str, headers: Sequence[str]) -> Any:
    if row is None:
        return None
    return row[headers.index(header)]


def _write_merged_formula(
    worksheet,
    row: int,
    first: int,
    last: int,
    formula: str,
    cell_format,
    cached: Any = "",
) -> None:
    worksheet.merge_range(row, first, row, last, "", cell_format)
    worksheet.write_formula(
        row,
        first,
        formula,
        cell_format,
        _formula_cache(cached),
    )


def _cache_map(rows: Sequence[Sequence[Any]]) -> dict[str, tuple[Any, ...]]:
    return {str(row[0]): tuple(row) for row in rows}


_LOOKUP_DATASETS: dict[str, tuple[str, Sequence[str]]] = {
    "tblPcsLob": ("PCS_LOB_DATA", PCS_LOB_SCORECARD_HEADERS),
    "tblPcsAgent": ("PCS_AGENT_DATA", PCS_AGENT_SCORECARD_HEADERS),
    "tblPcsDaily": ("PCS_DAILY_DATA", PCS_DAILY_SCORECARD_HEADERS),
    "tblPcsCoachingView": ("PCS_COACH_DATA", PCS_COACHING_HEADERS),
}


def _table_lookup(table: str, header: str, key_expression: str) -> str:
    """Return a lookup that survives Power Query table replacement.

    Desktop Excel must delete each starter ListObject before it can create the
    Power Query destination. A structured reference such as
    ``tblPcsLob[PCS]`` is rewritten to ``#REF!`` at deletion time. The stable
    workbook names below point to sheet ranges instead, so installation cannot
    mutate presentation formulas.
    """

    try:
        range_name, headers = _LOOKUP_DATASETS[table]
        column = headers.index(header) + 1
    except (KeyError, ValueError) as exc:
        raise ValueError(f"Unsupported PCS lookup {table}[{header}]") from exc
    return (
        f'=IFERROR(INDEX({range_name},MATCH({key_expression},'
        f'INDEX({range_name},0,1),0),{column}),"")'
    )


def _selected_key(*, sheet: str | None = None) -> str:
    prefix = f"'{sheet}'!" if sheet else ""
    return (
        f'{prefix}$A$3&"|"&SUBSTITUTE({prefix}$H$3,"|","/")&"|"&'
        f'SUBSTITUTE({prefix}$O$3,"|","/")&"|"&'
        f'SUBSTITUTE({prefix}$V$3,"|","/")'
    )


def _add_filter_names(workbook) -> None:
    keys = "'_PCS_FILTERS'!$A:$A"
    values = "'_PCS_FILTERS'!$C:$C"

    def group_range(group: str) -> str:
        first = f"MATCH({group},{keys},0)"
        return (
            f'=INDEX({values},{first}):'
            f'INDEX({values},{first}+COUNTIF({keys},{group})-1)'
        )

    # Whole-column INDEX ranges do not drift when desktop Excel replaces the
    # starter Power Query ListObject. The Windows installer recreates and
    # validates these same four names after every live query refresh.
    workbook.define_name(
        "PCS_PERIOD_LIST",
        group_range('"PERIOD"'),
    )
    workbook.define_name(
        "PCS_LOB_LIST",
        group_range('"LOB"'),
    )
    team_group = '"TL|"&SUBSTITUTE(OVERVIEW!$H$3,"|","/")'
    workbook.define_name(
        "PCS_TL_ACTIVE",
        group_range(team_group),
    )
    agent_group = (
        '"AGENT|"&SUBSTITUTE(OVERVIEW!$H$3,"|","/")&"|"&'
        'SUBSTITUTE(OVERVIEW!$O$3,"|","/")'
    )
    workbook.define_name(
        "PCS_AGENT_ACTIVE",
        group_range(agent_group),
    )
    # Dynamic sheet-backed ranges deliberately do not reference ListObject
    # names. Power Query installation replaces the starter tables; these names
    # and every dependent formula remain valid throughout that operation.
    for name, sheet, width in (
        ("PCS_LOB_DATA", "_PCS_LOB", len(PCS_LOB_SCORECARD_HEADERS)),
        ("PCS_AGENT_DATA", "_PCS_AGENT", len(PCS_AGENT_SCORECARD_HEADERS)),
        ("PCS_DAILY_DATA", "_PCS_DAILY", len(PCS_DAILY_SCORECARD_HEADERS)),
        ("PCS_COACH_DATA", "_PCS_COACH", len(PCS_COACHING_HEADERS)),
    ):
        workbook.define_name(
            name,
            f'=OFFSET(\'{sheet}\'!$A$5,0,0,'
            f'MAX(1,COUNTA(\'{sheet}\'!$A:$A)-3),{width})',
        )


def _add_overview(
    report: ExcelReport,
    filter_rows: Sequence[Sequence[Any]],
    lob_rows: Sequence[Sequence[Any]],
    agent_rows: Sequence[Sequence[Any]],
    daily_rows: Sequence[Sequence[Any]],
) -> None:
    workbook = report.workbook
    worksheet = workbook.add_worksheet("OVERVIEW")
    formats = make_v2_formats(workbook)
    configure_v2_cell_canvas(worksheet, formats, zoom=82)
    worksheet.hide_row_col_headers()
    worksheet.set_tab_color(COLORS["gold"])
    worksheet.freeze_panes(3, 0)
    write_v2_header(
        worksheet, formats, "PCS PERFORMANCE & COACHING",
        status="LIVE FILTERS", status_kind="LIVE",
    )

    defaults = ("Current MTD", "All", "All", "All")
    filters = (
        ("PERIOD", defaults[0], "=PCS_PERIOD_LIST"),
        ("LOB", defaults[1], "=PCS_LOB_LIST"),
        ("TEAM LEADER", defaults[2], "=PCS_TL_ACTIVE"),
        ("AGENT", defaults[3], "=PCS_AGENT_ACTIVE"),
    )
    for block, (label, value, source) in enumerate(filters):
        first = block * 7
        worksheet.merge_range(1, first, 1, first + 6, label, formats.filter_label)
        worksheet.merge_range(2, first, 2, first + 6, value, formats.filter_value)
        worksheet.data_validation(2, first, 2, first, {
            "validate": "list", "source": source,
            "input_title": label,
            "input_message": "Choose a governed value; reset child filters to All after changing a parent.",
            "error_title": "Invalid filter",
            "error_message": "Choose a value from the dropdown list.",
        })
    worksheet.set_row_pixels(1, 20)
    worksheet.set_row_pixels(2, 32)

    selected = _selected_key()
    default_key = "|".join(defaults)
    lob_map = _cache_map(lob_rows)
    agent_map = _cache_map(agent_rows)
    daily_map = _cache_map(daily_rows)
    kpi = lob_map.get(
        f"KPI|{default_key}", tuple(None for _ in PCS_LOB_SCORECARD_HEADERS),
    )
    card_specs = (
        ("SELECTED PCS", "PCS", "decimal"),
        ("PARTICIPATION", "Participation", "percent"),
        ("PRIOR COMPARABLE", "Prior PCS", "decimal"),
        ("CHANGE", "Change", "decimal"),
    )
    write_v2_kpis(worksheet, formats, tuple(
        (
            label,
            _table_lookup("tblPcsLob", header, f'"KPI|"&{selected}'),
            kind,
            _cached(kpi, header, PCS_LOB_SCORECARD_HEADERS),
        )
        for label, header, kind in card_specs
    ))

    lob_chart = workbook.add_chart({"type": "bar"})
    lob_chart.add_series({
        "name": "Selected", "categories": "='_PCS_CALC'!$A$2:$A$13",
        "values": "='_PCS_CALC'!$B$2:$B$13",
        "fill": {"color": COLORS["teal"]}, "border": {"none": True},
        "data_labels": {"value": True, "num_format": "0.00"},
    })
    lob_chart.add_series({
        "name": "Prior comparable", "categories": "='_PCS_CALC'!$A$2:$A$13",
        "values": "='_PCS_CALC'!$C$2:$C$13",
        "fill": {"color": COLORS["muted"]}, "border": {"none": True},
    })
    style_v2_chart(lob_chart, title="PCS BY LOB · SELECTED SCOPE", kind="bar")
    lob_chart.set_x_axis({"min": 1, "max": 5, "major_unit": 1})
    lob_chart.set_y_axis({"reverse": True})

    trend_chart = workbook.add_chart({"type": "line"})
    trend_chart.add_series({
        "name": "Daily PCS", "categories": "='_PCS_CALC'!$D$2:$D$32",
        "values": "='_PCS_CALC'!$E$2:$E$32",
        "line": {"color": COLORS["teal"], "width": 2.25},
        "marker": {"type": "circle", "size": 5,
                   "fill": {"color": COLORS["teal"]},
                   "border": {"color": COLORS["teal"]}},
    })
    style_v2_chart(trend_chart, title="DAILY PCS · SELECTED PERIOD")
    trend_chart.set_y_axis({"min": 1, "max": 5, "major_unit": 1})
    trend_chart.set_x_axis({"date_axis": True, "num_format": "d-mmm"})
    worksheet.insert_chart(7, 0, lob_chart)
    worksheet.insert_chart(7, 14, trend_chart)

    worksheet.merge_range(
        ACTION_SECTION_ROW, 0, ACTION_SECTION_ROW, 27,
        "LOB PERFORMANCE · RESPONDS TO ALL FOUR DROPDOWNS", formats.section,
    )
    lob_headers = (
        "LOB", "VALID Q1", "SELECTED PCS", "PRIOR PCS",
        "CHANGE", "PARTICIPATION", "COACHING DUE",
    )
    lob_source_headers = (
        "LOB", "Valid Q1", "PCS", "Prior PCS", "Change",
        "Participation", "Coaching Due",
    )
    lob_formats = (
        formats.table_text, formats.table_integer, formats.table_decimal,
        formats.table_decimal, formats.table_decimal, formats.table_percent,
        formats.table_integer,
    )
    for column, header in enumerate(lob_headers):
        first = column * 4
        worksheet.merge_range(
            ACTION_HEADER_ROW, first, ACTION_HEADER_ROW, first + 3,
            header, formats.table_header,
        )
    for offset in range(12):
        rank = offset + 1
        row_index = ACTION_FIRST_ROW + offset
        cached = lob_map.get(
            f"LOB|{default_key}|{rank}",
            tuple(None for _ in PCS_LOB_SCORECARD_HEADERS),
        )
        key = f'"LOB|"&{selected}&"|{rank}"'
        for column, (header, cell_format) in enumerate(zip(lob_source_headers, lob_formats)):
            first = column * 4
            _write_merged_formula(
                worksheet, row_index, first, first + 3,
                _table_lookup("tblPcsLob", header, key), cell_format,
                _cached(cached, header, PCS_LOB_SCORECARD_HEADERS),
            )
        worksheet.set_row_pixels(row_index, 26)

    agent_section_row = ACTION_FIRST_ROW + 14
    agent_header_row = agent_section_row + 1
    worksheet.merge_range(
        agent_section_row, 0, agent_section_row, 27,
        "AGENT PERFORMANCE · CASCADING LOB → TEAM LEADER → AGENT", formats.section,
    )
    agent_headers = (
        ("AGENT", 0, 4), ("AGENT ID", 5, 7), ("LOB", 8, 10),
        ("TEAM LEADER", 11, 15), ("SELECTED PCS", 16, 18),
        ("PRIOR PCS", 19, 21), ("CHANGE", 22, 24),
        ("PARTICIPATION", 25, 27),
    )
    agent_source_headers = (
        "Agent", "Agent ID", "LOB", "Team Leader", "PCS",
        "Prior PCS", "Change", "Participation",
    )
    agent_formats = (
        formats.table_text, formats.table_text, formats.table_text,
        formats.table_text, formats.table_decimal, formats.table_decimal,
        formats.table_decimal, formats.table_percent,
    )
    for header, first, last in agent_headers:
        worksheet.merge_range(
            agent_header_row, first, agent_header_row, last,
            header, formats.table_header,
        )
    agent_capacity = max(50, min(1000, sum(
        1 for row in filter_rows if row[0] == "AGENT|All|All"
    ) - 1))
    for offset in range(agent_capacity):
        rank = offset + 1
        row_index = agent_header_row + 1 + offset
        cached = agent_map.get(
            f"AGENT|{default_key}|{rank}",
            tuple(None for _ in PCS_AGENT_SCORECARD_HEADERS),
        )
        key = f'"AGENT|"&{selected}&"|{rank}"'
        for (_label, first, last), header, cell_format in zip(
            agent_headers, agent_source_headers, agent_formats,
        ):
            _write_merged_formula(
                worksheet, row_index, first, last,
                _table_lookup("tblPcsAgent", header, key), cell_format,
                _cached(cached, header, PCS_AGENT_SCORECARD_HEADERS),
            )
        worksheet.set_row_pixels(row_index, 25)
    worksheet.set_footer(
        "&LPrepared by Anass ASSRI | WFM&CPCS operational tracker&RPage &P of &N",
    )


def _add_calc_sheet(
    report: ExcelReport,
    lob_rows: Sequence[Sequence[Any]],
    daily_rows: Sequence[Sequence[Any]],
) -> None:
    worksheet = report.workbook.add_worksheet("_PCS_CALC")
    headers = ("LOB", "PCS", "Prior PCS", "Date", "Daily PCS", "Daily Participation")
    for column, header in enumerate(headers):
        worksheet.write(0, column, header, report.header)
    selected = _selected_key(sheet="OVERVIEW")
    default_key = "Current MTD|All|All|All"
    lob_map = _cache_map(lob_rows)
    daily_map = _cache_map(daily_rows)
    for offset in range(12):
        rank = offset + 1
        key = f'"LOB|"&{selected}&"|{rank}"'
        cached = lob_map.get(
            f"LOB|{default_key}|{rank}",
            tuple(None for _ in PCS_LOB_SCORECARD_HEADERS),
        )
        for column, header in enumerate(("LOB", "PCS", "Prior PCS")):
            worksheet.write_formula(
                offset + 1, column,
                _table_lookup("tblPcsLob", header, key),
                None, _formula_cache(_cached(cached, header, PCS_LOB_SCORECARD_HEADERS)),
            )
    for offset in range(31):
        rank = offset + 1
        key = f'"DAILY|"&{selected}&"|{rank}"'
        cached = daily_map.get(
            f"DAILY|{default_key}|{rank}",
            tuple(None for _ in PCS_DAILY_SCORECARD_HEADERS),
        )
        for column, header, cell_format in (
            (3, "Date", report.date), (4, "PCS", report.decimal),
            (5, "Participation", report.percent),
        ):
            worksheet.write_formula(
                offset + 1, column,
                _table_lookup("tblPcsDaily", header, key),
                cell_format,
                _formula_cache(_cached(cached, header, PCS_DAILY_SCORECARD_HEADERS)),
            )
    worksheet.hide()


def _add_coaching(
    report: ExcelReport,
    coaching_rows: Sequence[Sequence[Any]],
    actions: Sequence[Sequence[Any]],
) -> None:
    worksheet = report.workbook.add_worksheet("COACHING")
    worksheet.hide_gridlines(2)
    worksheet.set_zoom(78)
    worksheet.set_tab_color(COLORS["gold"])
    worksheet.freeze_panes(4, 0)
    worksheet.merge_range("A1:R1", "PCS COACHING", report.title)
    worksheet.merge_range("A2:B2", "PERIOD", report.subtitle)
    worksheet.merge_range("C2:E2", "Current MTD", report.editable)
    worksheet.merge_range("F2:G2", "LOB", report.subtitle)
    worksheet.merge_range("H2:J2", "All", report.editable)
    worksheet.merge_range(
        "L2:R2", "Blue table = permanent human-owned action log", report.subtitle,
    )
    worksheet.data_validation("C2", {"validate": "list", "source": "=PCS_PERIOD_LIST"})
    worksheet.data_validation("H2", {"validate": "list", "source": "=PCS_LOB_LIST"})
    queue_headers = PCS_COACHING_HEADERS[2:12]
    queue_capacity = max(30, min(5000, max(
        (int(row[1] or 0) for row in coaching_rows), default=0,
    )))
    default_key = "COACH|Current MTD|All"
    cache = _cache_map(coaching_rows)
    for offset in range(queue_capacity):
        rank = offset + 1
        cached = cache.get(
            f"{default_key}|{rank}", tuple(None for _ in PCS_COACHING_HEADERS),
        )
        key = f'"COACH|"&$C$2&"|"&SUBSTITUTE($H$2,"|","/")&"|{rank}"'
        for column, header in enumerate(queue_headers):
            worksheet.write_formula(
                4 + offset, column,
                _table_lookup("tblPcsCoachingView", header, key),
                _cell_format(report, header, _cached(cached, header, PCS_COACHING_HEADERS)),
                _formula_cache(_cached(cached, header, PCS_COACHING_HEADERS)),
            )
    worksheet.add_table(3, 0, 3 + queue_capacity, len(queue_headers) - 1, {
        "name": "tblCoachingQueue", "style": "Table Style Light 9",
        "columns": [{"header": header, "header_format": report.header} for header in queue_headers],
    })
    padded_actions = [tuple(row) for row in actions]
    padded_actions.extend(
        tuple(None for _ in COACHING_ACTION_HEADERS)
        for _ in range(max(25, 100 - len(padded_actions)))
    )
    _write_table(
        report, worksheet, first_row=3, first_column=11,
        table_name="tblCoachingActions", headers=COACHING_ACTION_HEADERS,
        rows=padded_actions, editable_headers=set(COACHING_ACTION_HEADERS),
    )
    worksheet.data_validation(4, 13, 5003, 13, {
        "validate": "list",
        "source": ["Pending", "Planned", "Completed", "Not required"],
    })
    worksheet.conditional_format(4, 11, 5003, 11, {
        "type": "duplicate", "format": report.error,
    })
    for column, width in enumerate((12, 13, 20, 24, 12, 10, 11, 24, 34, 40)):
        worksheet.set_column(column, column, width)
    worksheet.set_column(10, 10, 3)
    for column, width in enumerate((40, 24, 16, 20, 16, 16, 34), 11):
        worksheet.set_column(column, column, width)
    worksheet.set_footer(
        "&LPrepared by Anass ASSRI | WFM&CPCS coaching&RPage &P of &N",
    )


def _add_setup(report: ExcelReport, config: Config) -> None:
    folder = config.feed / "PCS"
    worksheet = report.add_table_sheet(
        "SETUP",
        "PCS CONNECTION SETUP",
        "Install once with WFMHub; normal updates replace CSV feeds, then Excel Data > Refresh All reloads the tables.",
        ["Setting", "Value", "Why it exists"],
        [
            ("Power Query Installed", "NO", "Set automatically after all six query tables refresh"),
            ("Connection Mode", "LOCAL", "Use the locally synced WFMHub feed folder"),
            ("Connection Owner", "Anass ASSRI", "One owner controls connection changes"),
            ("Local Feed Folder", str(folder), "Fixed clean CSV folder outside the workbook"),
            ("SharePoint Site URL", "https://company.sharepoint.com/sites/WFM", "Reserved for a later SharePoint-feed migration"),
            ("SharePoint Feed Folder", "/Shared Documents/WFMHub/Feed/PCS/", "Reserved SharePoint folder fragment"),
            ("Filter Script", str(folder / "POWER_QUERY_PCS_FILTERS_LOCAL.txt"), "Cascading Period, LOB, Team Leader and Agent lists"),
            ("LOB Script", str(folder / "POWER_QUERY_PCS_LOB_LOCAL.txt"), "Overview LOB scorecard"),
            ("Agent Script", str(folder / "POWER_QUERY_PCS_AGENT_LOCAL.txt"), "Overview agent scorecard"),
            ("Daily Script", str(folder / "POWER_QUERY_PCS_DAILY_LOCAL.txt"), "Overview daily trend"),
            ("Performance Script", str(folder / "POWER_QUERY_PCS_RESULTS_LOCAL.txt"), "Native-filter and slicer-ready performance table"),
            ("Coaching Script", str(folder / "POWER_QUERY_COACHING_QUEUE_LOCAL.txt"), "Period/LOB coaching view cache"),
            ("Workbook Last Refreshed", "Never", "Written after desktop Excel finishes Refresh All"),
            ("Last Installer Result", "Not run", "Latest desktop Excel connection result"),
            ("Template Version", PCS_TRACKER_VERSION, "Controls the one-time action-preserving migration"),
        ],
        editable_headers={"Value"},
    )
    worksheet.set_column("A:A", 28)
    worksheet.set_column("B:B", 78)
    worksheet.set_column("C:C", 58)


def _add_help(report: ExcelReport) -> None:
    worksheet = report.add_table_sheet(
        "HELP",
        "PCS — SIMPLE OPERATING MODEL",
        "Power Query is transport only. Python/SQLite calculate every governed result before Excel reads the CSVs.",
        ["Step", "What to do", "Where", "Result"],
        [
            (1, "Run Update PCS data", "WFMHub", "SQLite and the fixed clean CSV feeds are refreshed"),
            (2, "Run Install/repair Power Query once after this upgrade", "PCS menu", "The permanent workbook is connected to the CSV feeds"),
            (3, "Open PCS Live Tracker and choose Data > Refresh All", "Excel", "All six lightweight query caches reload"),
            (4, "Choose Period, then LOB, Team Leader and Agent", "OVERVIEW", "Cards, charts and both tables update together"),
            (5, "After changing a parent, reset child filters to All", "OVERVIEW", "Every selection remains valid and easy to understand"),
            (6, "Use Period and LOB", "COACHING", "The exact low-score calls appear with Call ID and Coaching Key"),
            (7, "Copy the key and call ID into the blue action table", "COACHING", "Quality records status, coach, dates and comment permanently"),
            (8, "Use table filters or add slicers for detailed checks", "PERFORMANCE", "Inspect governed period, LOB, team and agent rows"),
        ],
    )
    worksheet.set_column("A:A", 10)
    worksheet.set_column("B:B", 62)
    worksheet.set_column("C:C", 24)
    worksheet.set_column("D:D", 64)


def _add_audit(report: ExcelReport, generated: datetime) -> None:
    worksheet = report.add_table_sheet(
        "_AUDIT",
        "PCS TRACKER CONTRACT",
        "Technical support information.",
        ["Field", "Value"],
        [
            ("Tracker version", PCS_TRACKER_VERSION),
            ("Design", f"{REPORT_DESIGN_ID} {REPORT_DESIGN_VERSION}"),
            ("Created", generated),
            ("Workbook grain", "Precomputed selection caches and low-score calls"),
            ("Raw data", "External fixed CSV feeds only; no raw-data worksheet"),
            ("Update method", "Power Query from six fixed CSV feeds"),
            ("KPI method", "Python/SQLite ratio of additive sums"),
            ("Human-owned table", "COACHING!tblCoachingActions"),
        ],
    )
    worksheet.hide()


def ensure_pcs_tracker(
    config: Config,
    *,
    filter_rows: Sequence[Sequence[Any]] = (),
    lob_rows: Sequence[Sequence[Any]] = (),
    agent_rows: Sequence[Sequence[Any]] = (),
    daily_rows: Sequence[Sequence[Any]] = (),
    result_rows: Sequence[Sequence[Any]] = (),
    coaching_rows: Sequence[Sequence[Any]] = (),
) -> Path:
    """Create or explicitly migrate the permanent action-preserving tracker."""

    target = tracker_path(config)
    if target.is_file() and _tracker_contract_version(target) == PCS_TRACKER_VERSION:
        return target
    migrated_actions = _legacy_actions(config)
    target.parent.mkdir(parents=True, exist_ok=True)
    partial = target.with_name(f"{target.stem}.partial{target.suffix}")
    generated = datetime.now()
    report = ExcelReport(partial)
    report.workbook.set_calc_mode("auto")
    report.workbook.set_custom_property(
        "WFMHub Report Design", f"{REPORT_DESIGN_ID} {REPORT_DESIGN_VERSION}",
    )
    report.workbook.set_custom_property(
        "WFMHub PCS Live Tracker", PCS_TRACKER_VERSION,
    )
    try:
        _add_filter_names(report.workbook)
        _add_overview(report, filter_rows, lob_rows, agent_rows, daily_rows)
        _query_sheet(
            report,
            name="PERFORMANCE",
            title="PCS PERFORMANCE",
            subtitle="Lightweight Power Query output · use native table filters or add slicers · rates are ratio-of-sums results from the Hub",
            table_name="tblPcsPerformance",
            headers=PCS_RESULTS_HEADERS,
            rows=result_rows,
        )
        _add_coaching(report, coaching_rows, migrated_actions)
        _add_setup(report, config)
        _add_help(report)
        _query_sheet(
            report,
            name="_PCS_FILTERS",
            title="PCS FILTER STAGING",
            subtitle="Power Query destination for governed cascading dropdown lists.",
            table_name="tblPcsFilters",
            headers=PCS_FILTER_HEADERS,
            rows=filter_rows,
            hidden=True,
        )
        _query_sheet(
            report,
            name="_PCS_LOB",
            title="PCS LOB STAGING",
            subtitle="Power Query destination for the lightweight Overview LOB scorecard.",
            table_name="tblPcsLob",
            headers=PCS_LOB_SCORECARD_HEADERS,
            rows=lob_rows,
            hidden=True,
        )
        _query_sheet(
            report,
            name="_PCS_AGENT",
            title="PCS AGENT STAGING",
            subtitle="Power Query destination for the lightweight Overview agent scorecard.",
            table_name="tblPcsAgent",
            headers=PCS_AGENT_SCORECARD_HEADERS,
            rows=agent_rows,
            hidden=True,
        )
        _query_sheet(
            report,
            name="_PCS_DAILY",
            title="PCS DAILY STAGING",
            subtitle="Power Query destination for the lightweight Overview trend.",
            table_name="tblPcsDaily",
            headers=PCS_DAILY_SCORECARD_HEADERS,
            rows=daily_rows,
            hidden=True,
        )
        _query_sheet(
            report,
            name="_PCS_COACH",
            title="PCS COACHING STAGING",
            subtitle="Power Query destination for the period and LOB coaching cache.",
            table_name="tblPcsCoachingView",
            headers=PCS_COACHING_HEADERS,
            rows=coaching_rows,
            hidden=True,
        )
        _add_calc_sheet(report, lob_rows, daily_rows)
        _add_audit(report, generated)
        report.close()
        if target.is_file():
            archive_dir = target.parent / "Archive" / generated.strftime("%Y-%m-%d")
            archive_dir.mkdir(parents=True, exist_ok=True)
            archived = archive_dir / (
                f"{target.stem}_pre_{PCS_TRACKER_VERSION.replace('.', '_')}_"
                f"{generated:%Y%m%d_%H%M%S_%f}{target.suffix}"
            )
            shutil.copy2(target, archived)
        partial.replace(target)
    except PermissionError as exc:
        partial.unlink(missing_ok=True)
        raise RuntimeError(
            "Close PCS Live Tracker.xlsx and wait for OneDrive to finish syncing. "
            "This upgrade must replace the old tracker once; existing coaching "
            "actions were read before the replacement was attempted."
        ) from exc
    except Exception:
        partial.unlink(missing_ok=True)
        raise
    return target


def build_pcs_live_tracker(
    conn: DatabaseConnection,
    config: Config,
    start: date,
    end: date,
    output: Path | None = None,
) -> Path:
    """Publish fixed CSV feeds and ensure the permanent tracker exists."""

    del output
    first, last = conn.execute(
        "SELECT min(business_date), max(business_date) FROM mart.agent_pcs_day",
    ).fetchone()
    available_start = _as_date(first) or start
    available_end = _as_date(last) or end
    publish_pcs_feeds(conn, config, available_start, available_end)
    target = tracker_path(config)
    if target.is_file() and _tracker_contract_version(target) == PCS_TRACKER_VERSION:
        return target
    generated = datetime.now()
    filter_rows, lob_rows, agent_rows, daily_rows = pcs_dashboard_cache_rows(
        conn, available_end, generated,
    )
    metric_catalog = load_metric_catalog(config.home, config.metric_catalog)
    method = metric_catalog.method_for("pcs_average", available_end, {})
    minimum_sample = int(method.minimum_sample) if method is not None else 1
    result_rows = pcs_result_rows(
        conn, available_end, minimum_sample, generated,
    )
    coaching_rows = pcs_coaching_cache_rows(
        conn, config, available_end, generated,
    )
    return ensure_pcs_tracker(
        config,
        filter_rows=filter_rows,
        lob_rows=lob_rows,
        agent_rows=agent_rows,
        daily_rows=daily_rows,
        result_rows=result_rows,
        coaching_rows=coaching_rows,
    )
