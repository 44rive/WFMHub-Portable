"""Inspect and operate the permanent PCS workbook on Windows Excel.

The Python report generator can create a valid ``.xlsx`` file, but Power Query's
binary mashup parts are owned by desktop Excel.  The packaged Windows helper uses
Excel automation to install the two governed queries, refresh them, and save the
same collaborative workbook.
"""

from __future__ import annotations

import csv
import os
import shutil
import subprocess
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from zipfile import BadZipFile, ZipFile

from openpyxl import load_workbook

from .config import Config


PCS_TEMPLATE_VERSION = "2026.09.20"


class PCSExcelError(RuntimeError):
    """Raised when desktop Excel cannot install or refresh the PCS queries."""


@dataclass(frozen=True)
class PCSTrackerState:
    path: Path
    exists: bool
    template_version: str | None = None
    setup_state: str | None = None
    connection_mode: str | None = None
    workbook_data_through: str | None = None
    workbook_feed_refreshed_at: str | None = None
    workbook_refreshed_at: str | None = None
    feed_data_through: str | None = None
    feed_refreshed_at: str | None = None
    query_parts: int = 0
    has_connections: bool = False
    needs_excel_refresh: bool = False
    problem: str | None = None

    @property
    def current_template(self) -> bool:
        return self.template_version == PCS_TEMPLATE_VERSION

    @property
    def queries_installed(self) -> bool:
        return (
            self.setup_state == "YES"
            and self.has_connections
            and self.query_parts >= 2
        )


def _text(value: object) -> str | None:
    if value in (None, ""):
        return None
    if isinstance(value, datetime):
        return value.isoformat(sep=" ", timespec="seconds")
    return str(value).strip() or None


def _manifest_values(path: Path) -> dict[str, str]:
    if not path.is_file():
        return {}
    with path.open("r", encoding="utf-8-sig", newline="") as handle:
        return {
            str(row.get("Item") or "").strip(): str(row.get("Value") or "").strip()
            for row in csv.DictReader(handle)
            if row.get("Item")
        }


def _setup_values(workbook) -> dict[str, str]:
    if "SETUP" not in workbook.sheetnames:
        return {}
    sheet = workbook["SETUP"]
    return {
        str(row[0]).strip(): str(row[1]).strip()
        for row in sheet.iter_rows(min_row=5, values_only=True)
        if len(row) >= 2 and row[0] not in (None, "") and row[1] not in (None, "")
    }


def _pcs_data_freshness(workbook) -> tuple[str | None, str | None]:
    if "PCS_DATA" not in workbook.sheetnames:
        return None, None
    sheet = workbook["PCS_DATA"]
    headers = {
        str(cell.value).strip(): cell.column
        for cell in sheet[4]
        if cell.value not in (None, "")
    }
    through_col = headers.get("Data Through")
    refreshed_col = headers.get("Feed Refreshed At")
    through: list[str] = []
    refreshed: list[str] = []
    for row in sheet.iter_rows(min_row=5, values_only=True):
        if through_col and through_col <= len(row):
            value = _text(row[through_col - 1])
            if value:
                through.append(value)
        if refreshed_col and refreshed_col <= len(row):
            value = _text(row[refreshed_col - 1])
            if value:
                refreshed.append(value)
    return max(through, default=None), max(refreshed, default=None)


def inspect_pcs_tracker(path: Path, feed_folder: Path | None = None) -> PCSTrackerState:
    """Return truthful template, connection, and freshness state without editing."""

    path = path.resolve()
    manifest = _manifest_values(
        (feed_folder / "PCS_MANIFEST_CURRENT.csv") if feed_folder else Path()
    ) if feed_folder else {}
    feed_through = manifest.get("Data through")
    feed_refreshed = manifest.get("Last refreshed")
    if not path.is_file():
        return PCSTrackerState(
            path=path, exists=False, feed_data_through=feed_through,
            feed_refreshed_at=feed_refreshed,
        )
    try:
        with ZipFile(path) as archive:
            names = archive.namelist()
            has_connections = "xl/connections.xml" in names
            query_parts = sum(name.startswith("xl/queryTables/queryTable") for name in names)
        workbook = load_workbook(path, read_only=True, data_only=False, keep_links=False)
        try:
            setup = _setup_values(workbook)
            workbook_through, workbook_feed_refreshed = _pcs_data_freshness(workbook)
        finally:
            workbook.close()
    except (BadZipFile, OSError, ValueError) as exc:
        return PCSTrackerState(
            path=path, exists=True, feed_data_through=feed_through,
            feed_refreshed_at=feed_refreshed, problem=str(exc),
        )
    needs_refresh = bool(
        feed_refreshed and (
            not workbook_feed_refreshed or feed_refreshed > workbook_feed_refreshed
        )
    )
    return PCSTrackerState(
        path=path,
        exists=True,
        template_version=setup.get("Template Version"),
        setup_state=setup.get("Power Query Installed", "NO").upper(),
        connection_mode=setup.get("Connection Mode"),
        workbook_data_through=workbook_through,
        workbook_feed_refreshed_at=workbook_feed_refreshed,
        workbook_refreshed_at=setup.get("Workbook Last Refreshed"),
        feed_data_through=feed_through,
        feed_refreshed_at=feed_refreshed,
        query_parts=query_parts,
        has_connections=has_connections,
        needs_excel_refresh=needs_refresh,
    )


def _helper_path(config: Config) -> Path:
    packaged = config.home / "_system" / "scripts" / "Install-PCSWorkbook.ps1"
    if packaged.is_file():
        return packaged
    development = config.home / "packaging" / "windows" / "Install-PCSWorkbook.ps1"
    if development.is_file():
        return development
    raise PCSExcelError("The packaged PCS Excel helper is missing")


def _powershell() -> str:
    executable = shutil.which("powershell.exe") or shutil.which("powershell")
    if executable:
        return executable
    system_root = os.environ.get("SystemRoot")
    if system_root:
        candidate = Path(system_root) / "System32" / "WindowsPowerShell" / "v1.0" / "powershell.exe"
        if candidate.is_file():
            return str(candidate)
    raise PCSExcelError("Windows PowerShell was not found; desktop Excel setup cannot run")


def run_pcs_excel_action(
    config: Config,
    workbook: Path,
    action: str,
    mode: str = "LOCAL",
    *,
    open_after: bool = False,
) -> str:
    """Run the reviewed Excel COM helper on Windows and return its status line."""

    if os.name != "nt":
        raise PCSExcelError(
            "PCS Power Query installation and workbook refresh require Windows desktop Excel"
        )
    normalized_action = action.strip().title()
    if normalized_action not in {"Install", "Refresh"}:
        raise PCSExcelError(f"Unknown PCS Excel action: {action}")
    normalized_mode = mode.strip().upper()
    if normalized_mode not in {"LOCAL", "SHAREPOINT"}:
        raise PCSExcelError("PCS connection mode must be LOCAL or SHAREPOINT")
    folder = config.feed / "PCS"
    suffix = normalized_mode
    command = [
        _powershell(), "-NoLogo", "-NoProfile", "-NonInteractive",
        "-ExecutionPolicy", "Bypass", "-File", str(_helper_path(config)),
        "-Action", normalized_action,
        "-Mode", normalized_mode,
        "-WorkbookPath", str(workbook.resolve()),
        "-FeedFolder", str(folder.resolve()),
        "-DataQueryPath", str(folder / f"POWER_QUERY_PCS_DATA_{suffix}.txt"),
        "-CoachingQueryPath", str(folder / f"POWER_QUERY_COACHING_QUEUE_{suffix}.txt"),
    ]
    if open_after:
        command.append("-OpenAfter")
    result = subprocess.run(command, capture_output=True, text=True, check=False)
    output = "\n".join(
        value.strip() for value in (result.stdout, result.stderr) if value.strip()
    )
    if result.returncode:
        raise PCSExcelError(output or f"Excel helper stopped with code {result.returncode}")
    return output or f"PCS Excel {normalized_action.lower()} completed"


def open_pcs_tracker(path: Path) -> None:
    """Open the tracker in the user's associated desktop application."""

    if os.name != "nt":
        raise PCSExcelError("Opening the PCS tracker automatically is available on Windows only")
    if not path.is_file():
        raise PCSExcelError(f"PCS tracker does not exist: {path}")
    os.startfile(path)  # type: ignore[attr-defined]
