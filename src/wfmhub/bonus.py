"""Bonus Matrix import, calculation and standard report product."""

from __future__ import annotations

import hashlib
import re
from calendar import monthrange
from dataclasses import dataclass
from datetime import date, datetime
from pathlib import Path
from typing import Any
from xml.etree import ElementTree
from zipfile import BadZipFile, ZipFile

from openpyxl import load_workbook

from .config import Config
from .database import DatabaseConnection
from .decision_products import _atomic_book, _audit_rows, _finish, _ratio
from .excel_layout import V2ChartSpec, render_v2_dashboard
from .bonus_source import (
    _cached_result_total,
    _bonus_period,
    _bonus_value,
    _calculate_bonus,
    _clean,
    _number,
    _policy_map,
    _read_bonus_source,
)
from .template_reports import KpiCard
from .reports import COLORS
from .report_packs import report_current_path


BONUS_IMPORT_VERSION = "2026.09.1-org-fields"
BONUS_TRACKER_VERSION = "1.1.0"
BONUS_INPUT_ROWS = 1000

# Exact six populations and thresholds in the uploaded Bonus Matrix v1.2.
# These are editable workbook seeds, never a second payout authority.
_BONUS_POPULATIONS = (
    ("Ford GER", 1000, .20, .10),
    ("Ford Dutch", 450, .50, .40),
    ("OEM FR", 450, .20, .10),
    ("RSA FR", 450, .50, .40),
    ("RSA NL", 450, .50, .40),
    ("RSA VL", 450, .50, .40),
)
DEFAULT_BONUS_KPIS = tuple(
    rule
    for population, aht_target, participation_tier_1, participation_tier_2
    in _BONUS_POPULATIONS
    for rule in (
        (population, "AHT", "L", .15, aht_target, 0, 0),
        (population, "Productivity", "H", .15, 7, 0, 0),
        (population, "PCS Score", "H", .20, 4.3, .15, 4.1),
        (population, "PCS % (Participation)", "H", .15, participation_tier_1, .10, participation_tier_2),
        (population, "QM", "H", .20, .9, .15, .8),
        (population, "Abs%", "L", .15, .05, 0, 0),
        (population, "Extra Bonus (PCS Score)", "H", .30, 4.5, .20, 4.3),
    )
)


def _bonus_tracker_version(path: Path) -> str | None:
    try:
        with ZipFile(path) as archive:
            try:
                root = ElementTree.fromstring(archive.read("docProps/custom.xml"))
                for prop in root:
                    if prop.attrib.get("name") == "WFMHub Bonus Tracker":
                        return next((child.text for child in prop), None)
            except KeyError:
                pass
            # Some spreadsheet writers discard custom properties on save.
            # The fixed 1,000-row input and result tables are an independent
            # structural marker, so a manually maintained tracker stays stable.
            tables = {}
            for name in archive.namelist():
                if name.startswith("xl/tables/") and name.endswith(".xml"):
                    table = ElementTree.fromstring(archive.read(name))
                    tables[table.attrib.get("displayName") or table.attrib.get("name")] = table.attrib.get("ref", "")
            for required in ("tblRawData", "tblResults"):
                match = re.search(r"(\d+)$", tables.get(required, ""))
                if match is None or int(match.group(1)) < 4 + BONUS_INPUT_ROWS:
                    return None
            config_end = re.search(r"(\d+)$", tables.get("tblKpiConfig", ""))
            if config_end is None or int(config_end.group(1)) < 4 + len(DEFAULT_BONUS_KPIS):
                return None
            return BONUS_TRACKER_VERSION
    except (BadZipFile, KeyError, OSError, ElementTree.ParseError):
        return None


def _existing_bonus_inputs(path: Path) -> dict[str, list[tuple[Any, ...]]]:
    """Preserve manually entered cells when replacing an earlier bonus file."""

    if not path.is_file():
        return {}
    workbook = load_workbook(path, read_only=True, data_only=False, keep_links=False)
    try:
        output: dict[str, list[tuple[Any, ...]]] = {}
        for sheet_name, width, key_index in (
            ("Policy_Decisions", 7, 0),
            ("KPI_Config", 7, 0),
            ("Raw_Data", 21, 0),
        ):
            if sheet_name not in workbook.sheetnames:
                continue
            sheet = workbook[sheet_name]
            values = []
            for row in sheet.iter_rows(min_row=5, values_only=True):
                cells = tuple(None if isinstance(value, str) and value.startswith("=") else value
                              for value in row[:width])
                if cells[key_index] not in (None, ""):
                    values.append(cells)
            output[sheet_name] = values
        return output
    finally:
        workbook.close()


@dataclass(frozen=True)
class BonusImportResult:
    import_id: str
    period: str
    agents: int
    rules: int
    policies: int
    unchanged: bool


def _text(value: Any) -> str | None:
    cleaned = _clean(value)
    return cleaned or None


def import_bonus_matrix(conn: DatabaseConnection, source: Path) -> BonusImportResult:
    """Read Bonus Matrix v1.2 without editing it and activate one period version."""

    source = source.resolve()
    if not source.exists():
        raise FileNotFoundError(source)
    content = source.read_bytes()
    source_sha = hashlib.sha256(content).hexdigest()
    import_id = source_sha
    existing = conn.execute(
        "SELECT period, agent_rows, rule_rows, policy_rows, active, source_cached_total, import_version FROM raw.bonus_import WHERE import_id=?",
        [import_id],
    ).fetchone()
    if existing and bool(existing[4]) and existing[6] == BONUS_IMPORT_VERSION:
        raw_count = conn.execute(
            "SELECT count(*) FROM raw.bonus_agent_month WHERE import_id=?", [import_id],
        ).fetchone()[0]
        rule_count = conn.execute(
            "SELECT count(*) FROM raw.bonus_kpi_rule WHERE import_id=?", [import_id],
        ).fetchone()[0]
        policy_count = conn.execute(
            "SELECT count(*) FROM raw.bonus_policy WHERE import_id=?", [import_id],
        ).fetchone()[0]
        mart_count = conn.execute(
            "SELECT count(*) FROM mart.bonus_agent_month WHERE import_id=?", [import_id],
        ).fetchone()[0]
        if (
            raw_count == existing[1]
            and mart_count == existing[1]
            and rule_count == existing[2]
            and policy_count == existing[3]
        ):
            return BonusImportResult(import_id, existing[0], existing[1], existing[2], existing[3], True)

    records, rules, policies = _read_bonus_source(source)
    period = _bonus_period(records)
    if not records:
        raise ValueError("Bonus source contains no populated Raw_Data agent rows")
    policy_values = _policy_map(policies)
    policies_ready = bool(policies) and all(
        len(row) > 5 and _clean(row[5]).casefold() == "validated" for row in policies
    )
    results = [_calculate_bonus(record, rules, policy_values, policies_ready) for record in records]
    source_cached_total, source_cached_rows = _cached_result_total(source)

    agent_sql = """INSERT INTO raw.bonus_agent_month (
        import_id, source_row, period, agent_id, agent_name, population, aht,
        productivity, pcs_score, pcs_participation, qm, absence_rate,
        voc_detractors, currency, monthly_fixed_salary, target_bonus_rate,
        reference_bonus_override, eligible_days, scheduled_days,
        employment_status, data_status, notes, team_leader, ops_manager
    ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)"""
    rule_map = {(rule.population, rule.kpi): rule for rule in rules}
    actual_names = (
        ("AHT", ("AHT",)),
        ("Productivity", ("Productivity",)),
        ("PCS Score", ("PCS Score",)),
        ("PCS % (Participation)", ("PCS % Participation", "PCS % (Participation)")),
        ("QM", ("QM",)),
        ("Abs%", ("Abs%",)),
        ("Extra Bonus (PCS Score)", ("PCS Score",)),
    )

    conn.execute("SAVEPOINT import_bonus_matrix")
    try:
        conn.execute("UPDATE raw.bonus_import SET active=false WHERE period=?", [period])
        conn.execute(
            """INSERT OR REPLACE INTO raw.bonus_import
               (import_id, source_path, file_name, source_sha256, imported_at, period,
                active, agent_rows, rule_rows, policy_rows, source_cached_total,
                source_cached_rows, import_version)
               VALUES (?, ?, ?, ?, ?, ?, true, ?, ?, ?, ?, ?, ?)""",
            [import_id, str(source), source.name, source_sha, datetime.now(), period,
             len(records), len(rules), len(policies), source_cached_total,
             source_cached_rows, BONUS_IMPORT_VERSION],
        )
        conn.execute("DELETE FROM raw.bonus_agent_month WHERE import_id=?", [import_id])
        conn.execute("DELETE FROM raw.bonus_kpi_rule WHERE import_id=?", [import_id])
        conn.execute("DELETE FROM raw.bonus_policy WHERE import_id=?", [import_id])

        for row_number, record in enumerate(records, 1):
            conn.execute(agent_sql, [
                import_id, row_number, _text(_bonus_value(record, "Period")) or period,
                _text(_bonus_value(record, "Agent ID")), _text(_bonus_value(record, "Agent Name")),
                _text(_bonus_value(record, "Population")), _number(_bonus_value(record, "AHT")),
                _number(_bonus_value(record, "Productivity")), _number(_bonus_value(record, "PCS Score")),
                _number(_bonus_value(record, "PCS % Participation", "PCS % (Participation)")),
                _number(_bonus_value(record, "QM")), _number(_bonus_value(record, "Abs%")),
                int(_number(_bonus_value(record, "VOC Detractor Count")) or 0),
                _text(_bonus_value(record, "Currency")) or "MAD",
                _number(_bonus_value(record, "Monthly Fixed Salary")),
                _number(_bonus_value(record, "Target Bonus Rate")),
                _number(_bonus_value(record, "Reference Bonus Override")),
                _number(_bonus_value(record, "Eligible Days")),
                _number(_bonus_value(record, "Scheduled Days")),
                _text(_bonus_value(record, "Employment Status")),
                _text(_bonus_value(record, "Data Status")),
                _text(_bonus_value(record, "Notes")),
                _text(_bonus_value(record, "Team Lead", "Team Leader")),
                _text(_bonus_value(record, "Ops Manager", "Operations Manager")),
            ])
        for rule in rules:
            conn.execute(
                "INSERT INTO raw.bonus_kpi_rule VALUES (?, ?, ?, ?, ?, ?, ?, ?)",
                [import_id, rule.population, rule.kpi, rule.direction, rule.tier1_bonus,
                 rule.tier1_target, rule.tier2_bonus, rule.tier2_target],
            )
        for policy in policies:
            values = list(policy[:7]) + [None] * (7 - len(policy[:7]))
            conn.execute(
                "INSERT INTO raw.bonus_policy VALUES (?, ?, ?, ?, ?, ?, ?, ?)",
                [import_id, *[_text(value) for value in values[:7]]],
            )

        conn.execute("DELETE FROM mart.bonus_agent_month WHERE period=?", [period])
        conn.execute("DELETE FROM mart.bonus_kpi_result WHERE period=?", [period])
        for record, result in zip(records, results):
            agent_id = _text(_bonus_value(record, "Agent ID"))
            agent_name = _text(_bonus_value(record, "Agent Name"))
            population = _text(_bonus_value(record, "Population")) or ""
            earned = list(result.earned) + [None] * (7 - len(result.earned))
            conn.execute(
                """INSERT INTO mart.bonus_agent_month (
                   period, agent_id, agent_name, population, core_ready, eligibility,
                   aht_earned, productivity_earned, pcs_earned, participation_earned,
                   qm_earned, absence_earned, extra_pcs_earned, gross_achievement,
                   voc_malus, final_achievement, reference_bonus, proration,
                   scenario_payout, released_payout, release_status, data_issue,
                   import_id, source_sha256, team_leader, ops_manager
                ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)""",
                [period, agent_id, agent_name, population, result.core_ready, result.eligibility,
                 *earned[:7], result.gross, result.malus, result.final, result.reference,
                 result.proration, result.scenario, result.release, result.status,
                 result.issue, import_id, source_sha,
                 _text(_bonus_value(record, "Team Lead", "Team Leader")),
                 _text(_bonus_value(record, "Ops Manager", "Operations Manager"))],
            )
            for index, (kpi, aliases) in enumerate(actual_names):
                rule = rule_map.get((population, kpi))
                if rule is None:
                    continue
                conn.execute(
                    "INSERT INTO mart.bonus_kpi_result VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)",
                    [period, agent_id, agent_name, population, kpi,
                     _number(_bonus_value(record, *aliases)), earned[index], rule.direction,
                     rule.tier1_target, rule.tier2_target, import_id],
                )
        conn.execute("RELEASE SAVEPOINT import_bonus_matrix")
    except Exception:
        conn.execute("ROLLBACK TO SAVEPOINT import_bonus_matrix")
        conn.execute("RELEASE SAVEPOINT import_bonus_matrix")
        raise
    return BonusImportResult(import_id, period, len(records), len(rules), len(policies), False)


def _selected_period(conn: DatabaseConnection, end: date) -> str:
    wanted = end.strftime("%Y-%m")
    exists = conn.execute("SELECT 1 FROM mart.bonus_agent_month WHERE period=? LIMIT 1", [wanted]).fetchone()
    if exists:
        return wanted
    latest = conn.execute("SELECT max(period) FROM mart.bonus_agent_month").fetchone()[0]
    return latest or wanted


def _bonus_earned_formula(row: int, raw_column: str, kpi: str) -> str:
    """Return the v1.2 tier-one-then-tier-two award formula."""

    lookup = f'tblKpiConfig[Population]&"|"&tblKpiConfig[KPI]'
    wanted = f'$C{row}&"|{kpi}"'
    return (
        f'=IF($A{row}="","",LET(actual,Raw_Data!{raw_column}{row},'
        f'direction,XLOOKUP({wanted},{lookup},tblKpiConfig[Direction],""),'
        f'bonus1,XLOOKUP({wanted},{lookup},tblKpiConfig[Tier 1 Bonus %],0),'
        f'target1,XLOOKUP({wanted},{lookup},tblKpiConfig[Tier 1 Target],0),'
        f'bonus2,XLOOKUP({wanted},{lookup},tblKpiConfig[Tier 2 Bonus %],0),'
        f'target2,XLOOKUP({wanted},{lookup},tblKpiConfig[Tier 2 Target],0),'
        'IF(actual="",0,IF(direction="L",IF(actual<=target1,bonus1,'
        'IF(AND(bonus2>0,actual<=target2),bonus2,0)),IF(actual>=target1,bonus1,'
        'IF(AND(bonus2>0,actual>=target2),bonus2,0))))))'
    )


def _bonus_control_rows(period: str) -> list[tuple[Any, ...]]:
    rows = [
        ("Monetary inputs", "Missing currency", '=COUNTIFS(tblRawData[Agent ID],"<>",tblRawData[Currency],"")', 0, None, "WFM / HR", "Currency is required", None),
        ("Monetary inputs", "Currency different from MAD", '=COUNTIFS(tblRawData[Agent ID],"<>",tblRawData[Currency],"<>MAD")', 0, None, "WFM / HR", "Confirm conversion policy", None),
        ("Monetary inputs", "Missing salary/rate and override", '=SUMPRODUCT(--(tblRawData[Agent ID]<>""),--(tblRawData[Reference Bonus Override]=""),--((tblRawData[Monthly Fixed Salary]="")+(tblRawData[Target Bonus Rate]="")>0))', 0, None, "WFM / HR", "Reference amount cannot be calculated", None),
        ("Proration", "Missing or zero Scheduled Days", '=SUMPRODUCT(--(tblRawData[Agent ID]<>""),--((tblRawData[Scheduled Days]="")+(tblRawData[Scheduled Days]<=0)>0))', 0, None, "WFM / HR", "Proration denominator", None),
        ("Proration", "Eligible Days above Scheduled Days", '=SUMPRODUCT(--(tblRawData[Agent ID]<>""),--(tblRawData[Eligible Days]>tblRawData[Scheduled Days]))', 0, None, "WFM / HR", "Eligible days are capped at scheduled days", None),
        ("Period", "Period differs from Dashboard selection", '=COUNTIFS(tblRawData[Agent ID],"<>",tblRawData[Period],"<>"&Dashboard!$C$2)', 0, None, "WFM", "All rows should use the selected month", None),
        ("Configuration", "Population missing from KPI configuration", '=SUMPRODUCT(--(tblRawData[Agent ID]<>""),--(COUNTIF(tblKpiConfig[Population],tblRawData[Population])=0))', 0, None, "WFM / HR", "Add the population before release", None),
        ("Data", "Duplicate Agent ID and Period", '=SUMPRODUCT((tblRawData[Agent ID]<>"")*(COUNTIFS(tblRawData[Agent ID],tblRawData[Agent ID],tblRawData[Period],tblRawData[Period])>1))/2', 0, None, "WFM", "One agent row per month", None),
        ("Policy", "Policy decisions not validated", '=COUNTIF(tblPolicyDecisions[Status],"<>Validated")', 0, None, "WFM / HR", "Every policy decision needs an owner", None),
        ("Data status", "Rows requiring review", '=COUNTIFS(tblRawData[Agent ID],"<>",tblRawData[Data Status],"<>VALIDATED")', 0, None, "WFM / HR", "Confirm incomplete records", None),
        ("Roster", "Agents without Team Lead", '=COUNTIFS(tblRawData[Agent ID],"<>",tblRawData[Team Lead],"")', 0, None, "Operations", "Complete ownership mapping", None),
        ("Information", "Total loaded records", '=COUNTA(tblRawData[Agent ID])', "INFO", None, "WFM", "Rows available for calculation", None),
        ("Information", "Estimated final payout", '=SUM(tblResults[Final Payout])', "INFO", None, "WFM / HR", "Current workbook result", None),
    ]
    output = []
    for excel_row, values in enumerate(rows, 5):
        status = f'=IF(D{excel_row}="INFO","INFO",IF(C{excel_row}=D{excel_row},"OK","REVIEW"))'
        output.append((*values[:4], status, *values[5:]))
    return output


def build_bonus_performance_workbook(
    conn: DatabaseConnection,
    config: Config,
    start: date,
    end: date,
    output: Path | None = None,
) -> Path:
    """Create the permanent, formula-driven Bonus Matrix once."""

    target_path = output or report_current_path(config, "bonus")
    if output is None and target_path.is_file() and _bonus_tracker_version(target_path) == BONUS_TRACKER_VERSION:
        return target_path
    preserved = _existing_bonus_inputs(target_path)
    period = _selected_period(conn, end)
    period_start = date.fromisoformat(f"{period}-01")
    period_end = period_start.replace(day=monthrange(period_start.year, period_start.month)[1])
    book, partial, target = _atomic_book(
        config, "bonus", "BONUS MANAGEMENT", period_start, period_end, output,
    )
    book.report.workbook.set_calc_mode("auto")
    book.report.workbook.set_custom_property("WFMHub Bonus Tracker", BONUS_TRACKER_VERSION)
    import_row = conn.execute(
        """SELECT import_id, file_name, imported_at
           FROM raw.bonus_import WHERE period=? AND active=true
           ORDER BY imported_at DESC LIMIT 1""",
        [period],
    ).fetchone()
    import_id = import_row[0] if import_row else None

    policy_headers = [
        "policy", "selected_decision", "allowed_values", "formula_impact",
        "owner", "status", "comments",
    ]
    policy_rows = conn.execute(
        """SELECT policy, selected_value, allowed_values, formula_impact,
                  owner, status, comments
           FROM raw.bonus_policy WHERE import_id=? ORDER BY rowid""",
        [import_id],
    ).fetchall() if import_id else []
    if preserved.get("Policy_Decisions"):
        policy_rows = preserved["Policy_Decisions"]
    if not policy_rows:
        policy_rows = [
            ("Malus method", "Proportional", "Proportional | Percentage points", "Final achievement after VOC malus", "HR / Compensation", "To validate", None),
            ("Extra PCS treatment", "Additive", "Additive | Replacement", "How the extra PCS award combines with PCS", "Operations / HR", "To validate", None),
            ("Target bonus rate basis", "Monthly", "Monthly | Annual", "Annual rates are divided by 12", "HR / Payroll", "To validate", None),
            ("Proration denominator", "Calendar days", "Calendar days | Working days | Planned days", "Eligible Days / Scheduled Days", "HR / Payroll", "To validate", None),
            ("Absence treatment", "KPI only", "KPI only | Eligibility only | Both", "Avoid duplicate absence impact", "HR / Operations", "To validate", None),
            ("Achievement cap", 1.3, "0% to 200%", "Maximum gross achievement", "HR / Compensation", "To validate", None),
            ("Rounding", "Centime", "Centime | Whole MAD", "Final payout decimals", "Payroll", "To validate", None),
            ("Employee eligibility", "Active eligible population", "Policy reference required", "Joiners, leavers and exclusions", "HR", "To validate", None),
        ]
    policy_sheet = book.table(
        "Policy_Decisions", "Policy decisions",
        "Confirm the selected decision, owner and status before the monthly result is released.",
        policy_headers, policy_rows,
        editable_headers={"Selected Decision", "Owner", "Status", "Comments"},
    )
    status_column = policy_headers.index("status")
    if policy_rows:
        policy_sheet.data_validation(
            4, status_column, 3 + len(policy_rows), status_column,
            {"validate": "list", "source": ["Validated", "To validate", "Rejected"]},
        )

    control_headers = [
        "area", "check", "count_value", "expected", "status", "owner",
        "evidence", "comment",
    ]
    control_rows = _bonus_control_rows(period)
    controls = book.table(
        "Control_Checks", "Control checks",
        "Resolve REVIEW items before management or payroll use.",
        control_headers, control_rows,
        editable_headers={"Comment"},
    )
    controls.conditional_format(
        4, 4, 3 + len(control_rows), 4,
        {"type": "text", "criteria": "containing", "value": "REVIEW", "format": book.report.error},
    )

    kpi_headers = [
        "population", "kpi", "direction", "tier_1_bonus_percent",
        "tier_1_target", "tier_2_bonus_percent", "tier_2_target",
    ]
    kpi_rows = conn.execute(
        """SELECT population, kpi, direction, tier1_bonus, tier1_target,
                  tier2_bonus, tier2_target
           FROM raw.bonus_kpi_rule WHERE import_id=?
           ORDER BY population,
             CASE kpi WHEN 'AHT' THEN 1 WHEN 'Productivity' THEN 2
                      WHEN 'PCS Score' THEN 3
                      WHEN 'PCS % (Participation)' THEN 4 WHEN 'QM' THEN 5
                      WHEN 'Abs%' THEN 6 ELSE 7 END""",
        [import_id],
    ).fetchall() if import_id else []
    if preserved.get("KPI_Config"):
        kpi_rows = preserved["KPI_Config"]
        if len(kpi_rows) == 14 and {str(row[0]) for row in kpi_rows} == {"OEM", "RSA BNL"}:
            # The interim two-population seed did not match the uploaded v1.2
            # matrix. Replace only that exact shape; the previous workbook is
            # archived by publish_report so manual edits remain recoverable.
            kpi_rows = list(DEFAULT_BONUS_KPIS)
    if not kpi_rows:
        kpi_rows = list(DEFAULT_BONUS_KPIS)
    kpi_sheet = book.table(
        "KPI_Config", "Bonus configuration",
        "Population-specific targets and awards. Direction H means higher is better; L means lower is better.",
        kpi_headers, kpi_rows,
        editable_headers={
            "Direction", "Tier 1 Bonus %", "Tier 1 Target",
            "Tier 2 Bonus %", "Tier 2 Target",
        },
    )
    voc_rows = [(0, 0), (1, .10), (2, .20), (3, .50), (4, .75), (5, 1.0)]
    kpi_sheet.write(3, 8, "VOC Detractor Count", book.report.header)
    kpi_sheet.write(3, 9, "Malus Impact", book.report.header)
    for offset, values in enumerate(voc_rows, 4):
        kpi_sheet.write(offset, 8, values[0], book.report.integer)
        kpi_sheet.write(offset, 9, values[1], book.report.percent)
    kpi_sheet.add_table(3, 8, 3 + len(voc_rows), 9, {
        "name": "tblVocMalus", "style": "Table Style Light 9",
        "columns": [{"header": "VOC Detractor Count"}, {"header": "Malus Impact"}],
    })
    kpi_sheet.set_column(8, 9, 22)

    raw_headers = [
        "agent_id", "agent_name", "population", "period", "aht",
        "productivity", "pcs_score", "pcs_participation", "qm",
        "absence_percent", "voc_detractor_count", "notes", "currency",
        "monthly_fixed_salary", "target_bonus_rate",
        "reference_bonus_override", "eligible_days", "scheduled_days",
        "data_status", "team_lead", "ops_manager",
    ]
    raw_rows = conn.execute(
        """SELECT r.agent_id, r.agent_name, r.population, r.period, r.aht,
                  r.productivity, r.pcs_score, r.pcs_participation, r.qm,
                  r.absence_rate, r.voc_detractors, r.notes, r.currency,
                  r.monthly_fixed_salary, r.target_bonus_rate,
                  r.reference_bonus_override, r.eligible_days, r.scheduled_days,
                  r.data_status, coalesce(r.team_leader,d.team_leader),
                  coalesce(r.ops_manager,d.ops_manager)
           FROM raw.bonus_agent_month r
           LEFT JOIN core.dim_agent d ON d.agent_id=r.agent_id
           WHERE r.import_id=? ORDER BY r.population, r.agent_name""",
        [import_id],
    ).fetchall() if import_id else []
    if preserved.get("Raw_Data"):
        raw_rows = preserved["Raw_Data"]
    if len(raw_rows) > BONUS_INPUT_ROWS:
        raise ValueError(f"Bonus input has {len(raw_rows)} rows; capacity is {BONUS_INPUT_ROWS}")
    raw_rows = list(raw_rows) + [tuple(None for _ in raw_headers)] * (BONUS_INPUT_ROWS - len(raw_rows))
    raw_sheet = book.table(
        "Raw_Data", "Raw performance data",
        f"Monthly inputs for {period}. Paste or validate values here before reviewing Results.",
        raw_headers, raw_rows,
        editable_headers=set(raw_headers),
    )
    raw_sheet.set_column(0, 0, 18, book.report.workbook.add_format({"num_format": "@"}))
    raw_sheet.data_validation(4, 18, 3 + BONUS_INPUT_ROWS, 18, {
        "validate": "list", "source": ["TEMPLATE", "TO REVIEW", "VALIDATED"],
    })
    raw_sheet.set_column(13, 14, None, None, {"hidden": True})

    result_headers = [
        "Agent ID", "Agent Name", "LOB", "Period", "AHT Earned",
        "Productivity Earned", "PCS Earned", "Participation Earned",
        "QM Earned", "Abs Earned", "Extra PCS Earned", "Gross Achievement",
        "VOC Detractor Count", "Malus Impact", "Final Achievement", "Status",
        "Currency", "Monthly Fixed Salary", "Target Bonus Rate",
        "Reference Bonus Amount", "Proration Factor", "Prorated Bonus Base",
        "Gross Payout", "Malus Deduction", "Final Payout", "Team Lead",
        "Ops Manager",
    ]
    result_rows: list[tuple[Any, ...]] = []
    row_count = BONUS_INPUT_ROWS
    for offset in range(row_count):
        row = offset + 5
        result_rows.append((
            f'=IF(Raw_Data!A{row}="","",Raw_Data!A{row})',
            f'=IF(Raw_Data!B{row}="","",Raw_Data!B{row})',
            f'=IF(Raw_Data!C{row}="","",Raw_Data!C{row})',
            f'=IF(Raw_Data!D{row}="","",Raw_Data!D{row})',
            _bonus_earned_formula(row, "E", "AHT"),
            _bonus_earned_formula(row, "F", "Productivity"),
            _bonus_earned_formula(row, "G", "PCS Score"),
            _bonus_earned_formula(row, "H", "PCS % (Participation)"),
            _bonus_earned_formula(row, "I", "QM"),
            _bonus_earned_formula(row, "J", "Abs%"),
            _bonus_earned_formula(row, "G", "Extra Bonus (PCS Score)"),
            f'=IF($A{row}="","",MIN(XLOOKUP("Achievement cap",tblPolicyDecisions[Policy],tblPolicyDecisions[Selected Decision],1.3),IF(XLOOKUP("Extra PCS treatment",tblPolicyDecisions[Policy],tblPolicyDecisions[Selected Decision],"Additive")="Additive",SUM(E{row}:K{row}),SUM(E{row}:J{row})-G{row}+MAX(G{row},K{row}))))',
            f'=IF($A{row}="","",Raw_Data!K{row})',
            f'=IF($A{row}="","",XLOOKUP(M{row},tblVocMalus[VOC Detractor Count],tblVocMalus[Malus Impact],1,1))',
            f'=IF($A{row}="","",IF(XLOOKUP("Malus method",tblPolicyDecisions[Policy],tblPolicyDecisions[Selected Decision],"Proportional")="Percentage points",MAX(0,L{row}-N{row}),L{row}*(1-N{row})))',
            f'=IF($A{row}="","",IF(COUNTIF(tblPolicyDecisions[Status],"<>Validated")>0,"Policy review",IF(Raw_Data!S{row}<>"VALIDATED","Data review",IF(U{row}="","Proration review",IF(O{row}=0,"No payout",IF(N{row}>0,"Malus applied","Eligible"))))))',
            f'=IF($A{row}="","",Raw_Data!M{row})',
            f'=IF($A{row}="","",Raw_Data!N{row})',
            f'=IF($A{row}="","",Raw_Data!O{row})',
            f'=IF($A{row}="","",IF(Raw_Data!P{row}>0,Raw_Data!P{row},Raw_Data!N{row}*Raw_Data!O{row}/IF(XLOOKUP("Target bonus rate basis",tblPolicyDecisions[Policy],tblPolicyDecisions[Selected Decision],"Monthly")="Annual",12,1)))',
            f'=IF($A{row}="","",IF(OR(Raw_Data!R{row}="",Raw_Data!R{row}<=0),"",MIN(1,MAX(0,Raw_Data!Q{row}/Raw_Data!R{row}))))',
            f'=IF(OR($A{row}="",U{row}=""),"",T{row}*U{row})',
            f'=IF(OR($A{row}="",V{row}=""),"",V{row}*L{row})',
            f'=IF(OR($A{row}="",V{row}=""),"",V{row}*L{row}-V{row}*O{row})',
            f'=IF(OR($A{row}="",V{row}=""),"",ROUND(V{row}*O{row},IF(XLOOKUP("Rounding",tblPolicyDecisions[Policy],tblPolicyDecisions[Selected Decision],"Centime")="Whole MAD",0,2)))',
            f'=IF($A{row}="","",Raw_Data!T{row})',
            f'=IF($A{row}="","",Raw_Data!U{row})',
        ))
    results = book.table(
        "Results", "Bonus results",
        "Formula-driven v1.2 calculation. Filter Population, Team Lead or Status for review.",
        result_headers, result_rows,
    )
    results.conditional_format(
        4, 14, 3 + len(result_rows), 14,
        {"type": "3_color_scale", "min_color": COLORS["red_light"],
         "mid_color": COLORS["amber_light"], "max_color": COLORS["green_light"]},
    )

    kpi_analysis_headers = [
        "kpi", "population", "configured_agents", "average_actual",
        "average_earned_weight", "agents_earning", "attainment_rate",
    ]
    actual_column = {
        "AHT": "E", "Productivity": "F", "PCS Score": "G",
        "PCS % (Participation)": "H", "QM": "I", "Abs%": "J",
        "Extra Bonus (PCS Score)": "G",
    }
    earned_column = {
        "AHT": "E", "Productivity": "F", "PCS Score": "G",
        "PCS % (Participation)": "H", "QM": "I", "Abs%": "J",
        "Extra Bonus (PCS Score)": "K",
    }
    last = 4 + BONUS_INPUT_ROWS
    kpi_analysis_rows = []
    for index, rule in enumerate(kpi_rows):
        row = index + 5
        kpi = str(rule[1])
        actual = actual_column.get(kpi, "E")
        earned = earned_column.get(kpi, "E")
        kpi_analysis_rows.append((
            f"=KPI_Config!B{row}", f"=KPI_Config!A{row}",
            f'=COUNTIFS(Raw_Data!$C$5:$C${last},B{row},Raw_Data!$D$5:$D${last},Dashboard!$C$2)',
            f'=IFERROR(AVERAGEIFS(Raw_Data!${actual}$5:${actual}${last},Raw_Data!$C$5:$C${last},B{row},Raw_Data!$D$5:$D${last},Dashboard!$C$2),0)',
            f'=IFERROR(AVERAGEIFS(Results!${earned}$5:${earned}${last},Results!$C$5:$C${last},B{row},Results!$D$5:$D${last},Dashboard!$C$2),0)',
            f'=COUNTIFS(Results!$C$5:$C${last},B{row},Results!$D$5:$D${last},Dashboard!$C$2,Results!${earned}$5:${earned}${last},">0")',
            f'=IFERROR(F{row}/C{row},0)',
        ))
    kpi_analysis = book.table(
        "KPI_Analysis", "KPI analysis",
        "Coverage and target attainment by KPI and population.",
        kpi_analysis_headers,
        kpi_analysis_rows,
    )
    if kpi_analysis_rows:
        chart = book.report.workbook.add_chart({"type": "column"})
        chart.add_series({
            "name": "Attainment rate",
            "categories": ["KPI_Analysis", 4, 0, 3 + len(kpi_analysis_rows), 0],
            "values": ["KPI_Analysis", 4, 6, 3 + len(kpi_analysis_rows), 6],
            "fill": {"color": COLORS["teal"]}, "border": {"none": True},
        })
        chart.set_title({"name": "KPI attainment"})
        chart.set_y_axis({"num_format": "0%", "major_gridlines": {"visible": False}})
        chart.set_legend({"none": True})
        kpi_analysis.insert_chart("I5", chart, {"x_scale": 1.1, "y_scale": 1.0})

    tl_headers = [
        "team_lead", "population", "agents", "paid_agents", "payout_rate",
        "total_payout", "payout_share", "average_payout", "zero_payout",
        "review_items", "average_achievement",
    ]
    tl_rows = []
    for index in range(100):
        row = index + 5
        tl_rows.append((
            f'=IFERROR(INDEX(SORT(UNIQUE(FILTER(Raw_Data!$T$5:$T${last},'
            f'(Raw_Data!$T$5:$T${last}<>"")*(Raw_Data!$D$5:$D${last}=Dashboard!$C$2)))),'
            f'ROWS($A$5:A{row})),"")',
            "All",
            f'=IF(A{row}="","",COUNTIFS(Results!$Z$5:$Z${last},A{row},Results!$D$5:$D${last},Dashboard!$C$2))',
            f'=IF(A{row}="","",COUNTIFS(Results!$Z$5:$Z${last},A{row},Results!$D$5:$D${last},Dashboard!$C$2,Results!$Y$5:$Y${last},">0"))',
            f'=IFERROR(D{row}/C{row},0)',
            f'=IF(A{row}="","",SUMIFS(Results!$Y$5:$Y${last},Results!$Z$5:$Z${last},A{row},Results!$D$5:$D${last},Dashboard!$C$2))',
            f'=IFERROR(F{row}/SUMIFS(Results!$Y$5:$Y${last},Results!$D$5:$D${last},Dashboard!$C$2),0)',
            f'=IFERROR(F{row}/D{row},0)',
            f'=IF(A{row}="","",C{row}-D{row})',
            f'=IF(A{row}="","",COUNTIFS(Results!$Z$5:$Z${last},A{row},Results!$D$5:$D${last},Dashboard!$C$2,Results!$P$5:$P${last},"*review*"))',
            f'=IFERROR(AVERAGEIFS(Results!$O$5:$O${last},Results!$Z$5:$Z${last},A{row},Results!$D$5:$D${last},Dashboard!$C$2),0)',
        ))
    tl_sheet = book.table(
        "Team_Lead_Analysis", "Team Lead analysis",
        "Team ownership, payout distribution and review volume.",
        tl_headers, tl_rows,
    )
    if tl_rows:
        chart = book.report.workbook.add_chart({"type": "bar"})
        chart.add_series({
            "name": "Total payout",
            "categories": ["Team_Lead_Analysis", 4, 0, 23, 0],
            "values": ["Team_Lead_Analysis", 4, 5, 23, 5],
            "fill": {"color": COLORS["gold"]}, "border": {"none": True},
        })
        chart.set_title({"name": "Payout by Team Lead"})
        chart.set_legend({"none": True})
        chart.set_y_axis({"major_gridlines": {"visible": False}})
        tl_sheet.insert_chart("M5", chart, {"x_scale": 1.15, "y_scale": 1.1})

    input_period = next((str(row[3]) for row in raw_rows if row[0] not in (None, "") and row[3]), period)
    populations = list(dict.fromkeys(str(row[0]) for row in kpi_rows if row[0]))
    population_rows = []
    for index, population in enumerate(populations[:8]):
        row = index + 25
        population_rows.append((
            population,
            f'=COUNTIFS(Results!$C$5:$C${last},A{row},Results!$D$5:$D${last},$C$2)',
            f'=COUNTIFS(Results!$C$5:$C${last},A{row},Results!$D$5:$D${last},$C$2,Results!$Y$5:$Y${last},">0")',
            f'=IFERROR(C{row}/B{row},0)',
            f'=SUMIFS(Results!$Y$5:$Y${last},Results!$C$5:$C${last},A{row},Results!$D$5:$D${last},$C$2)',
            f'=IFERROR(AVERAGEIFS(Results!$Y$5:$Y${last},Results!$C$5:$C${last},A{row},Results!$D$5:$D${last},$C$2,Results!$Y$5:$Y${last},">0"),0)',
            f'=IFERROR(AVERAGEIFS(Results!$O$5:$O${last},Results!$C$5:$C${last},A{row},Results!$D$5:$D${last},$C$2),0)',
            f'=COUNTIFS(Results!$C$5:$C${last},A{row},Results!$D$5:$D${last},$C$2,Results!$P$5:$P${last},"*review*")',
        ))
    payout = f'=SUMIFS(Results!$Y$5:$Y${last},Results!$D$5:$D${last},$C$2)'
    paid = f'=COUNTIFS(Results!$D$5:$D${last},$C$2,Results!$Y$5:$Y${last},">0")'
    average_payout = f'=IFERROR({payout[1:]}/{paid[1:]},0)'
    review_items = f'=COUNTIFS(Results!$D$5:$D${last},$C$2,Results!$P$5:$P${last},"*review*")'
    status_text = "Fill Raw_Data; review KPI_Config and Policy_Decisions; results and charts recalculate in Excel."
    dashboard = book.report.workbook.add_worksheet("Dashboard")
    render_v2_dashboard(
        book.report.workbook, dashboard,
        title="BONUS MANAGEMENT",
        filters=(
            ("Period", input_period), ("Inputs", "Raw_Data"),
            ("Targets", "KPI_Config"), ("Currency", "MAD"),
        ),
        kpis=(
            ("Total payout", payout, "money"),
            ("Paid agents", paid, "integer"),
            ("Avg paid payout", average_payout, "money"),
            ("Review items", review_items, "integer"),
        ),
        left_chart=V2ChartSpec(
            "PAYOUT BY POPULATION", "column",
            tuple(str(row[0]) for row in population_rows),
            (("Total payout", tuple(row[4] for row in population_rows), COLORS["gold"]),),
        ),
        right_chart=V2ChartSpec(
            "KPI ATTAINMENT", "bar",
            tuple(f"{row[0]} / {row[1]}" for row in kpi_analysis_rows),
            (("Attainment rate", tuple(row[6] for row in kpi_analysis_rows), COLORS["teal"]),),
            "percent", 0, 1,
        ),
        action_title="Population payout summary",
        action_headers=("POPULATION", "AGENTS", "PAID AGENTS", "PAYOUT RATE", "TOTAL PAYOUT", "AVG ACHIEVEMENT", "REVIEW"),
        action_rows=tuple((row[0], row[1], row[2], row[3], row[4], row[6], row[7]) for row in population_rows),
        action_kinds=("text", "integer", "integer", "percent", "money", "percent", "alert"),
        status="BONUS WORKSPACE",
        status_kind="LIVE",
        status_note=status_text,
    )
    dashboard.data_validation("C2", {
        "validate": "custom", "value": '=AND(LEN(C2)=7,MID(C2,5,1)="-",ISNUMBER(VALUE(LEFT(C2,4))),ISNUMBER(VALUE(RIGHT(C2,2))))',
        "input_title": "Reporting month", "input_message": "Use YYYY-MM, for example 2026-09.",
    })
    dashboard.activate()
    if import_row:
        dashboard.write_comment("A1", f"Source: {import_row[1]}\nImported: {import_row[2]}", {
            "author": "Anass ASSRI",
        })
    return _finish(book, partial, target)
