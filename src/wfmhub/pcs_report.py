"""Python-only PCS snapshot and permanent coaching-log lifecycle.

PCS calculations belong to SQLite/Python.  The generated report is a disposable
snapshot containing final values and native charts; it never asks desktop Excel
to import, calculate, refresh, or save anything.  Human coaching actions live in
one separate workbook that WFMHub reads but never replaces after creation.
"""

from __future__ import annotations

import os
from collections import defaultdict
from datetime import date, datetime, timedelta
from pathlib import Path
from typing import Any, Mapping, Sequence

from openpyxl import load_workbook

from .config import Config
from .database import DatabaseConnection
from .design import COLORS, REPORT_DESIGN_ID, REPORT_DESIGN_VERSION
from .excel_layout import V2ChartSpec, render_v2_dashboard
from .metrics import load_metric_catalog
from .reports import ExcelReport, _query
from .rules import load_rulebook
from .shared_feeds import (
    PCS_AGENT_DAY_HEADERS,
    PCS_COACHING_HEADERS,
    PCS_RESULTS_HEADERS,
    pcs_lob_scorecard_rows,
    pcs_result_rows,
)
from .template_reports import DecisionWorkbook


PCS_COACHING_LOG_FILENAME = "PCS Coaching Log.xlsx"
PCS_REPORT_PREFIX = "PCS Operational Report - "
PCS_REPORT_GLOB = f"{PCS_REPORT_PREFIX}*.xlsx"
PCS_SNAPSHOT_VERSION = "2026.09.26"
PCS_SNAPSHOT_RESULTS_HEADERS = tuple(
    "Report Generated At" if header == "Feed Refreshed At" else header
    for header in PCS_RESULTS_HEADERS
)
PCS_SNAPSHOT_DATA_HEADERS = tuple(
    "Report Generated At" if header == "Feed Refreshed At" else header
    for header in PCS_AGENT_DAY_HEADERS
)

COACHING_LOG_HEADERS = (
    "Coaching Key", "Call ID", "LOB", "Team Leader", "Agent Selector",
    "Agent ID", "Agent", "Date", "Call Start", "Q1 Score",
    "Customer Comment", "Call Reference Number", "Language",
    "Coaching Status", "Coach", "Coaching Date", "Due Date",
    "Coaching Comment",
)
COACHING_QUEUE_HEADERS = (*COACHING_LOG_HEADERS[:13], "Priority", "Action Status",
                          "Coach", "Coaching Date", "Due Date", "Coaching Comment")
LOB_SUMMARY_HEADERS = (
    "LOB", "Current MTD PCS", "Prior MTD PCS", "Change",
    "Current MTD Participation", "Valid Q1", "Score <= 3", "Sample State",
)
DAILY_TREND_HEADERS = (
    "Date", "Current PCS", "Prior Comparable PCS", "Participation",
    "Valid Q1", "Score <= 3",
)


def coaching_log_path(config: Config) -> Path:
    return (config.reports / PCS_COACHING_LOG_FILENAME).resolve()


def latest_pcs_report(config: Config) -> Path | None:
    candidates = [
        path for path in config.reports.glob(PCS_REPORT_GLOB)
        if path.is_file() and ".partial" not in path.stem.casefold()
    ]
    return max(candidates, key=lambda path: path.stat().st_mtime, default=None)


def open_workbook(path: Path) -> None:
    if not path.is_file():
        raise FileNotFoundError(path)
    if os.name != "nt":
        raise RuntimeError(f"Open this workbook manually: {path}")
    os.startfile(path)  # type: ignore[attr-defined]


def _read_coaching_rows(path: Path) -> tuple[list[dict[str, Any]], set[str]]:
    """Read saved actions by Coaching Key without editing the workbook."""

    if not path.is_file():
        return [], set()
    workbook = load_workbook(path, read_only=True, data_only=True, keep_links=False)
    try:
        if "COACHING" not in workbook.sheetnames:
            return [], set()
        sheet = workbook["COACHING"]
        headers = {
            str(cell.value or "").strip(): cell.column
            for cell in sheet[4]
            if cell.value not in (None, "")
        }
        key_column = headers.get("Coaching Key")
        if key_column is None:
            return [], set()
        rows: list[dict[str, Any]] = []
        seen: set[str] = set()
        duplicates: set[str] = set()
        for values in sheet.iter_rows(min_row=5, values_only=True):
            key = values[key_column - 1] if key_column <= len(values) else None
            key_text = str(key or "").strip()
            if not key_text:
                continue
            if key_text in seen:
                duplicates.add(key_text)
                continue
            rows.append({
                header: values[column - 1] if column <= len(values) else None
                for header, column in headers.items()
                if header in COACHING_LOG_HEADERS
            })
            seen.add(key_text)
        return rows, duplicates
    finally:
        workbook.close()


def _queue_by_key(queue_rows: Sequence[Sequence[Any]]) -> dict[str, dict[str, Any]]:
    output: dict[str, dict[str, Any]] = {}
    for values in queue_rows:
        item = dict(zip(PCS_COACHING_HEADERS, values))
        key = str(item.get("Coaching Key") or "").strip()
        if key:
            output[key] = item
    return output


def _legacy_coaching_rows(config: Config) -> list[dict[str, Any]]:
    legacy = (config.reports / "PCS Operational Tracker.xlsx").resolve()
    try:
        rows, _duplicates = _read_coaching_rows(legacy)
        return rows
    except Exception:
        # The legacy workbook remains untouched. A locked or damaged copy must
        # never prevent creation of the new Python-only report architecture.
        return []


def ensure_coaching_log(
    config: Config,
    queue_rows: Sequence[Sequence[Any]] = (),
) -> Path:
    """Create the permanent coaching workbook once and never replace it."""

    target = coaching_log_path(config)
    if target.is_file():
        return target
    target.parent.mkdir(parents=True, exist_ok=True)
    queue = _queue_by_key(queue_rows)
    migrated = _legacy_coaching_rows(config)
    prepared: list[tuple[Any, ...]] = []
    for saved in migrated:
        key = str(saved.get("Coaching Key") or "").strip()
        source = queue.get(key, {})
        combined = {
            header: saved.get(header) if saved.get(header) not in (None, "")
            else source.get(header)
            for header in COACHING_LOG_HEADERS
        }
        combined["Coaching Key"] = key
        prepared.append(tuple(combined.get(header) for header in COACHING_LOG_HEADERS))

    # Ready-to-paste rows keep the first-use workflow obvious. The Hub reads
    # only keyed rows and never rewrites this file after the atomic creation.
    prepared.extend(
        tuple(None for _ in COACHING_LOG_HEADERS)
        for _ in range(max(100, 250 - len(prepared)))
    )
    partial = target.with_name(f"{target.stem}.partial{target.suffix}")
    report = ExcelReport(partial)
    report.workbook.set_custom_property(
        "WFMHub Report Design", f"{REPORT_DESIGN_ID} {REPORT_DESIGN_VERSION}",
    )
    report.workbook.set_custom_property("WFMHub PCS Coaching Log", PCS_SNAPSHOT_VERSION)
    try:
        sheet = report.add_table_sheet(
            "COACHING", "PCS COACHING LOG",
            "Paste columns A:M from the report's COACHING_QUEUE, then complete the blue action columns N:R. This file is permanent; WFMHub only reads it.",
            list(COACHING_LOG_HEADERS), prepared,
            editable_headers=set(COACHING_LOG_HEADERS),
        )
        sheet.freeze_panes(4, 0)
        sheet.data_validation("N5:N10004", {
            "validate": "list",
            "source": ["Pending", "Planned", "Completed", "Not required"],
        })
        sheet.conditional_format("A5:A10004", {
            "type": "duplicate", "format": report.error,
        })
        sheet.set_column("A:A", 34)
        sheet.set_column("B:B", 24)
        sheet.set_column("C:D", 20)
        sheet.set_column("E:E", 30)
        sheet.set_column("F:J", 18)
        sheet.set_column("K:K", 38)
        sheet.set_column("L:Q", 20)
        sheet.set_column("R:R", 42)
        sheet.set_footer("&LPrepared by Anass ASSRI | WFM&CPCS coaching actions&RPage &P of &N")
        report.close()
        partial.replace(target)
    except Exception:
        partial.unlink(missing_ok=True)
        raise
    return target


def _daily_trend_rows(
    conn: DatabaseConnection,
    latest: date,
) -> list[tuple[Any, ...]]:
    current_start = latest.replace(day=1)
    previous_end = current_start - timedelta(days=1)
    previous_start = previous_end.replace(day=1)
    cursor = conn.execute(
        """SELECT business_date, coalesce(sum(pcs_score_sum),0),
                  coalesce(sum(survey_responses),0),
                  coalesce(sum(pcs_participation_responses),0),
                  coalesce(sum(pcs_status_calls),0),
                  coalesce(sum(low_score_responses),0)
           FROM mart.agent_pcs_day
           WHERE business_date BETWEEN ? AND ?
           GROUP BY business_date ORDER BY business_date""",
        [previous_start, latest],
    )
    daily: dict[date, tuple[float, int, int, int, int]] = {}
    for raw_date, score_sum, valid, nonblank, eligible, low in cursor.fetchall():
        business_date = (
            raw_date.date() if isinstance(raw_date, datetime)
            else raw_date if isinstance(raw_date, date)
            else date.fromisoformat(str(raw_date)[:10])
        )
        daily[business_date] = (
            float(score_sum or 0), int(valid or 0), int(nonblank or 0),
            int(eligible or 0), int(low or 0),
        )
    rows = []
    for offset in range(latest.day):
        current_date = current_start + timedelta(days=offset)
        prior_date = previous_start + timedelta(days=offset)
        current = daily.get(current_date, (0.0, 0, 0, 0, 0))
        prior = daily.get(prior_date, (0.0, 0, 0, 0, 0)) if prior_date <= previous_end else (0.0, 0, 0, 0, 0)
        rows.append((
            current_date,
            current[0] / current[1] if current[1] else None,
            prior[0] / prior[1] if prior[1] else None,
            current[2] / current[3] if current[3] else None,
            current[1], current[4],
        ))
    return rows


def _coaching_action_map(
    rows: Sequence[Mapping[str, Any]],
) -> dict[str, Mapping[str, Any]]:
    return {
        str(row.get("Coaching Key") or "").strip(): row
        for row in rows
        if str(row.get("Coaching Key") or "").strip()
    }


def _team_action_rows(
    result_rows: Sequence[Sequence[Any]],
    queue_rows: Sequence[Sequence[Any]],
    actions: Mapping[str, Mapping[str, Any]],
    latest: date,
) -> list[tuple[Any, ...]]:
    indexes = {header: index for index, header in enumerate(PCS_RESULTS_HEADERS)}
    current_teams = {
        (str(row[indexes["LOB"]]), str(row[indexes["Team Leader"]])): row
        for row in result_rows
        if row[indexes["Period View"]] == "Current MTD"
        and row[indexes["Scope Level"]] == "TEAM"
    }
    prior_teams = {
        (str(row[indexes["LOB"]]), str(row[indexes["Team Leader"]])): row
        for row in result_rows
        if row[indexes["Period View"]] == "Previous MTD same days"
        and row[indexes["Scope Level"]] == "TEAM"
    }
    agents: dict[tuple[str, str], set[str]] = defaultdict(set)
    for row in result_rows:
        if row[indexes["Period View"]] != "Current MTD" or row[indexes["Scope Level"]] != "AGENT":
            continue
        agents[(str(row[indexes["LOB"]]), str(row[indexes["Team Leader"]]))].add(
            str(row[indexes["Agent ID"]]),
        )
    month_start = latest.replace(day=1)
    opportunities: dict[tuple[str, str], set[str]] = defaultdict(set)
    completed: dict[tuple[str, str], set[str]] = defaultdict(set)
    queue_indexes = {header: index for index, header in enumerate(PCS_COACHING_HEADERS)}
    for row in queue_rows:
        raw_date = row[queue_indexes["Date"]]
        call_date = (
            raw_date.date() if isinstance(raw_date, datetime)
            else raw_date if isinstance(raw_date, date)
            else date.fromisoformat(str(raw_date)[:10])
        )
        if not month_start <= call_date <= latest:
            continue
        team_key = (str(row[queue_indexes["LOB"]]), str(row[queue_indexes["Team Leader"]]))
        coaching_key = str(row[queue_indexes["Coaching Key"]])
        opportunities[team_key].add(coaching_key)
        action = actions.get(coaching_key, {})
        if str(action.get("Coaching Status") or "").strip().casefold() == "completed":
            completed[team_key].add(coaching_key)

    output = []
    for key, current in current_teams.items():
        prior = prior_teams.get(key)
        current_pcs = current[indexes["PCS Average"]]
        prior_pcs = prior[indexes["PCS Average"]] if prior else None
        change = (
            float(current_pcs) - float(prior_pcs)
            if current_pcs is not None and prior_pcs is not None else None
        )
        due = max(0, len(opportunities[key]) - len(completed[key]))
        output.append((
            key[1], len(agents[key]), current[indexes["Participation Rate"]],
            current_pcs, prior_pcs, change, due,
        ))
    output.sort(key=lambda row: (-int(row[6] or 0), row[3] is None, row[3] or 0, str(row[0])))
    return output


def _add_overview(
    book: DecisionWorkbook,
    status: str,
    status_note: str,
    latest: date,
    lob_rows: Sequence[Sequence[Any]],
    trend_rows: Sequence[Sequence[Any]],
    team_rows: Sequence[Sequence[Any]],
    target: float | None,
) -> None:
    worksheet = book.report.workbook.add_worksheet("OVERVIEW")
    all_row = next((row for row in lob_rows if str(row[0]) == "ALL"), (None,) * 19)
    display_status = "DATA FRESH" if status in {"LIVE", "FINAL"} else "CHECK DATA"
    lob_detail = [row for row in lob_rows if str(row[0]) != "ALL"]
    target_text = f" · TARGET {target:.2f}" if target is not None else ""
    render_v2_dashboard(
        book.report.workbook, worksheet,
        title="PCS OPERATIONS",
        filters=(
            ("Period", "Current MTD"),
            ("Data through", latest),
            ("Scope", "All active FTE"),
            ("Workbook", "Python snapshot"),
        ),
        kpis=(
            ("Current PCS", all_row[5], "decimal"),
            ("Participation", all_row[8], "percent"),
            ("Prior PCS", all_row[6], "decimal"),
            ("Change", all_row[7], "decimal"),
        ),
        left_chart=V2ChartSpec(
            f"PCS BY LOB{target_text}", "bar",
            [row[0] for row in lob_detail],
            (
                ("Current MTD", [row[5] for row in lob_detail], COLORS["teal"]),
                ("Prior comparable", [row[6] for row in lob_detail], COLORS["muted"]),
            ),
            minimum=1, maximum=5,
        ),
        right_chart=V2ChartSpec(
            "DAILY PCS TREND", "line",
            [row[0] for row in trend_rows],
            (
                ("Current month", [row[1] for row in trend_rows], COLORS["teal"]),
                ("Prior comparable", [row[2] for row in trend_rows], COLORS["muted"]),
            ),
            minimum=2, maximum=5,
        ),
        action_title="TEAM PERFORMANCE & COACHING",
        action_headers=(
            "TEAM LEADER", "PCS AGENTS", "PARTICIPATION", "CURRENT PCS",
            "PRIOR PCS", "CHANGE", "COACHING DUE",
        ),
        action_rows=team_rows,
        action_kinds=("text", "integer", "percent", "decimal", "decimal", "change", "alert"),
        status=display_status,
        status_kind=status,
        status_note=status_note,
    )


def _build_queue_rows(
    source_rows: Sequence[Sequence[Any]],
    actions: Mapping[str, Mapping[str, Any]],
    duplicates: set[str],
) -> list[tuple[Any, ...]]:
    output = []
    for source in source_rows:
        source_record = dict(zip(PCS_COACHING_HEADERS, source))
        key = str(source_record.get("Coaching Key") or "").strip()
        action = actions.get(key, {})
        status = "DUPLICATE KEY IN LOG" if key in duplicates else str(
            action.get("Coaching Status") or "Pending",
        )
        output.append((
            *(source_record.get(header) for header in COACHING_LOG_HEADERS[:13]),
            source_record.get("Priority"), status, action.get("Coach"),
            action.get("Coaching Date"), action.get("Due Date"),
            action.get("Coaching Comment"),
        ))
    return output


def build_pcs_snapshot_workbook(
    conn: DatabaseConnection,
    config: Config,
    start: date,
    end: date,
    output: Path | None = None,
) -> Path:
    """Generate one self-contained PCS workbook without opening Excel."""

    from .decision_products import _audit_rows, _pcs_coaching_rows, _source_state

    generated = datetime.now()
    target = (
        output.resolve() if output is not None else
        (config.reports / f"{PCS_REPORT_PREFIX}{generated:%Y-%m-%d %H%M%S}.xlsx").resolve()
    )
    if target.exists() and output is None:
        target = target.with_name(
            f"{PCS_REPORT_PREFIX}{generated:%Y-%m-%d %H%M%S_%f}.xlsx",
        )
    target.parent.mkdir(parents=True, exist_ok=True)
    partial = target.with_name(f"{target.stem}.partial{target.suffix}")

    latest_value = conn.execute("SELECT max(business_date) FROM mart.agent_pcs_day").fetchone()[0]
    if latest_value is None:
        raise RuntimeError("No PCS data exists. Add a valid Call-by-Call extract first.")
    latest = (
        latest_value.date() if isinstance(latest_value, datetime)
        else latest_value if isinstance(latest_value, date)
        else date.fromisoformat(str(latest_value)[:10])
    )
    metric_catalog = load_metric_catalog(config.home, config.metric_catalog)
    pcs_method = metric_catalog.method_for("pcs_average", latest, {})
    minimum_sample = int(pcs_method.minimum_sample) if pcs_method is not None else 1
    rulebook = load_rulebook(config.home, config.business_rules)
    month_index = latest.year * 12 + latest.month - 1
    first_index = month_index - (config.pcs_tracker.history_months - 1)
    data_start = min(start, date(first_index // 12, first_index % 12 + 1, 1))

    lob_rows = pcs_lob_scorecard_rows(conn, latest, minimum_sample, generated)
    result_rows = pcs_result_rows(conn, latest, minimum_sample, generated)
    trend_rows = _daily_trend_rows(conn, latest)
    action_headers, raw_actions = _pcs_coaching_rows(conn, config, data_start, latest)
    indexes = {header: index for index, header in enumerate(action_headers)}
    sources = {
        "LOB": "lob", "Team Leader": "team_leader", "Agent Selector": "agent_selector",
        "Agent": "agent_name", "Agent ID": "agent_id", "Priority": "priority",
        "Date": "business_date", "Call Start": "call_start", "Q1 Score": "q1_score",
        "Customer Comment": "customer_comment", "Call Reference Number": "call_reference_number",
        "Call ID": "call_id", "Language": "language", "Coaching Key": "coaching_key",
    }
    queue_source = [
        tuple(values[indexes[sources[header]]] for header in PCS_COACHING_HEADERS)
        for values in raw_actions
    ]
    log_path = ensure_coaching_log(config, queue_source)
    coaching_rows, duplicates = _read_coaching_rows(log_path)
    coaching = _coaching_action_map(coaching_rows)
    queue_rows = _build_queue_rows(queue_source, coaching, duplicates)
    team_rows = _team_action_rows(result_rows, queue_source, coaching, latest)
    status, status_note = _source_state(conn, ("fte", "calls"), latest)

    book = DecisionWorkbook(
        partial, config, "pcs", "PCS OPERATIONAL REPORT", data_start, latest, generated,
    )
    try:
        _add_overview(
            book, status, status_note, latest, lob_rows, trend_rows, team_rows,
            pcs_method.target if pcs_method is not None else None,
        )
        lob_summary = [
            (row[0], row[5], row[6], row[7], row[8], row[10], row[13], row[16])
            for row in lob_rows
        ]
        lob_sheet = book.table(
            "LOB_SUMMARY", "PCS LOB SUMMARY",
            "Current MTD and prior-month same-days results. Values are calculated by Python before the workbook is created.",
            list(LOB_SUMMARY_HEADERS), lob_summary,
        )
        lob_sheet.conditional_format("H5:H1004", {
            "type": "text", "criteria": "containing", "value": "LOW SAMPLE",
            "format": book.report.error,
        })
        book.table(
            "DAILY_TREND", "PCS DAILY TREND",
            "Current-month daily PCS against the matching day of the previous month.",
            list(DAILY_TREND_HEADERS), trend_rows,
        )
        result_sheet = book.table(
            "RESULTS", "PCS TEAM AND AGENT RESULTS",
            "Filter from left to right: Period View, Scope Level, LOB, Team Leader and Agent. Values are final; no recalculation is required.",
            list(PCS_SNAPSHOT_RESULTS_HEADERS), result_rows,
        )
        if pcs_method is not None and pcs_method.target is not None:
            result_sheet.conditional_format("K5:K100004", {
                "type": "cell", "criteria": "<", "value": pcs_method.target,
                "format": book.report.error,
            })
        queue_sheet = book.table(
            "COACHING_QUEUE", "PCS COACHING OPPORTUNITIES",
            f"Filter the queue, copy columns A:M into {PCS_COACHING_LOG_FILENAME}, then complete its blue action fields. Call ID opens the exact call.",
            list(COACHING_QUEUE_HEADERS), queue_rows,
        )
        queue_sheet.conditional_format("N5:N100004", {
            "type": "text", "criteria": "containing", "value": "HIGH",
            "format": book.report.error,
        })
        queue_sheet.conditional_format("O5:O100004", {
            "type": "text", "criteria": "containing", "value": "DUPLICATE",
            "format": book.report.error,
        })
        coaching_values = [
            tuple(row.get(header) for header in COACHING_LOG_HEADERS)
            for row in coaching_rows
        ]
        book.table(
            "COACHING", "PCS COACHING STATUS",
            f"Read-only snapshot from {PCS_COACHING_LOG_FILENAME}. Edit the separate permanent log, then build a fresh PCS report.",
            list(COACHING_LOG_HEADERS), coaching_values,
        )
        _headers, data_rows = _query(
            conn,
            """SELECT lob, team_leader,
                      coalesce(agent_name,'Agent') || ' [' || agent_id || ']',
                      agent_id, agent_name, business_date, ops_manager, language,
                      inbound_calls, pcs_status_calls, pcs_participation_responses,
                      survey_responses, pcs_score_sum, pcs_average,
                      pcs_participation_rate, low_score_responses,
                      top_box_responses, pcs_invalid_responses,
                      CASE WHEN survey_responses<? THEN 'LOW_SAMPLE' ELSE 'OK' END,
                      agent_id || '|' || business_date, ?, ?, ?, ?, ?, ?
               FROM mart.agent_pcs_day WHERE business_date BETWEEN ? AND ?
               ORDER BY business_date, lob, team_leader, agent_name""",
            [
                minimum_sample, latest, generated, rulebook.version, rulebook.sha256,
                metric_catalog.version, metric_catalog.sha256, data_start, latest,
            ],
        )
        book.table(
            "PCS_DATA", "PCS CLEAN AGENT-DAY DATA",
            "One row per in-scope agent and day. Use native filters or build optional pivots from this static table.",
            list(PCS_SNAPSHOT_DATA_HEADERS), data_rows,
        )
        book.table(
            "HELP", "PCS OPERATING GUIDE",
            "The report is generated entirely by Python. Excel is only the viewing and filtering tool.",
            ["Role", "Action", "Where", "Important"],
            [
                ("WFM", "Paste the latest untouched FTE and Call-by-Call extracts, then choose Build latest PCS report", "WFMHub", "No workbook needs to be closed"),
                ("WFM", "Open or send the newest timestamped report", "Reports", "Every report is self-contained"),
                ("Quality / TL", "Filter COACHING_QUEUE and copy columns A:M for the selected call", "Generated report", "Coaching Key and Call ID identify the exact call"),
                ("Quality / TL", "Paste A:M and complete status, coach, dates and comment", PCS_COACHING_LOG_FILENAME, "This is the only permanent editable file"),
                ("WFM", "Build again from the current database after coaching updates", "PCS menu", "Fast path: source extracts are not scanned again"),
            ],
        )
        book.definitions([
            ("PCS Average", "Sum of valid inbound Q1 scores / valid inbound Q1 responses", "Weighted LOB, team and agent result", "Never average agent averages"),
            ("PCS Participation", "Inbound raw Q1 nonblank / inbound PCSStatus=1", "Survey participation opportunity", "Invalid nonblank Q1 remains in the numerator"),
            ("Score <= 3", "Count of valid inbound Q1 responses at or below 3", "Coaching opportunity", "One exact call equals one Coaching Key"),
            ("Current MTD", "First day of the latest data month through the latest data date", "Current operating result", "Driven by the data-through date"),
            ("Previous MTD same days", "Previous month through the comparable day number", "Fair MTD comparison", "Month length is capped safely"),
            ("Python snapshot", "All cards, charts and tables contain final values", "No refresh dependency", "Excel performs no KPI calculation"),
        ])
        audit = _audit_rows(conn, config, "pcs", data_start, latest, (
            ("Workbook engine", "Python-only snapshot", "No Power Query, Excel automation, Data Model or dashboard formulas"),
            ("Snapshot version", PCS_SNAPSHOT_VERSION, "Generated workbook contract"),
            ("Coaching log", str(log_path), "Permanent human-owned workbook; read only by WFMHub"),
            ("Duplicate Coaching Keys", len(duplicates), "Duplicate keys do not increase completed coaching counts"),
        ))
        book.audit(audit)
        book.close()
        partial.replace(target)
    except Exception:
        partial.unlink(missing_ok=True)
        raise
    return target
