#!/usr/bin/env python3
"""Validate the FTE workbook boundary used by the PTO/Away Power Platform app."""

from __future__ import annotations

import argparse
import json
from collections import Counter
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.utils.cell import range_boundaries


TABLE_CONTRACTS = {
    "tblFTEAgents": {
        "sheet": "Agent",
        "headers": [
            "Client ID", "Status", "Name", "Team leader", "Ops Manager",
            "LOB", "Market", "Language", "Location", "City", "FTE",
            "End date if leaver",
        ],
    },
    "tblFTEPTO": {
        "sheet": "PTO",
        "headers": [
            "Client ID", "Name", "Start date", "End date", "Day coverage",
            "Start time", "End time", "PTO type", "Approval status", "Comment",
        ],
    },
    "tblFTEAway": {
        "sheet": "Away",
        "headers": [
            "Client ID", "Name", "Start date", "End date", "Away type",
            "Case status", "Comment",
        ],
    },
}


def _clean(value: Any) -> str:
    return str(value if value is not None else "").strip()


def inspect_fte_contract(path: Path) -> dict[str, Any]:
    result: dict[str, Any] = {
        "path": str(path),
        "ok": True,
        "errors": [],
        "warnings": [],
        "tables": {},
        "roster": {
            "populated_rows": 0,
            "blank_client_ids": 0,
            "numeric_client_ids": 0,
            "duplicate_client_ids": [],
            "leavers_without_end_date": 0,
        },
    }
    if not path.is_file():
        result["errors"].append("Workbook does not exist.")
        result["ok"] = False
        return result

    workbook = load_workbook(path, read_only=False, data_only=False, keep_links=False)
    try:
        for table_name, contract in TABLE_CONTRACTS.items():
            sheet_name = contract["sheet"]
            if sheet_name not in workbook.sheetnames:
                result["errors"].append(f"Required sheet {sheet_name!r} is missing.")
                continue
            sheet = workbook[sheet_name]
            if table_name not in sheet.tables:
                result["errors"].append(
                    f"Required table {table_name!r} is missing from {sheet_name!r}."
                )
                continue
            table = sheet.tables[table_name]
            min_col, min_row, max_col, max_row = range_boundaries(table.ref)
            headers = [
                sheet.cell(min_row, column).value
                for column in range(min_col, max_col + 1)
            ]
            result["tables"][table_name] = {
                "sheet": sheet_name,
                "reference": table.ref,
                "headers": headers,
            }
            if headers != contract["headers"]:
                result["errors"].append(
                    f"{table_name} headers changed. Expected {contract['headers']!r}; "
                    f"found {headers!r}."
                )

        if "Agent" in workbook.sheetnames and "tblFTEAgents" in workbook["Agent"].tables:
            sheet = workbook["Agent"]
            table = sheet.tables["tblFTEAgents"]
            min_col, min_row, max_col, max_row = range_boundaries(table.ref)
            headers = [
                _clean(sheet.cell(min_row, column).value)
                for column in range(min_col, max_col + 1)
            ]
            indexes = {header: offset + min_col for offset, header in enumerate(headers)}
            ids: list[str] = []
            for row_number in range(min_row + 1, max_row + 1):
                values = [
                    sheet.cell(row_number, column).value
                    for column in range(min_col, max_col + 1)
                ]
                if not any(value not in (None, "") for value in values):
                    continue
                result["roster"]["populated_rows"] += 1
                raw_id = sheet.cell(row_number, indexes["Client ID"]).value
                client_id = _clean(raw_id)
                if not client_id:
                    result["roster"]["blank_client_ids"] += 1
                else:
                    ids.append(client_id)
                    if isinstance(raw_id, (int, float)) and not isinstance(raw_id, bool):
                        result["roster"]["numeric_client_ids"] += 1
                status = _clean(sheet.cell(row_number, indexes["Status"]).value).upper()
                leave_date = sheet.cell(row_number, indexes["End date if leaver"]).value
                if status == "LEAVER" and leave_date in (None, ""):
                    result["roster"]["leavers_without_end_date"] += 1

            result["roster"]["duplicate_client_ids"] = sorted(
                client_id for client_id, count in Counter(ids).items() if count > 1
            )

        roster = result["roster"]
        if roster["blank_client_ids"]:
            result["warnings"].append(
                f"Roster has {roster['blank_client_ids']} populated rows without Client ID."
            )
        if roster["numeric_client_ids"]:
            result["warnings"].append(
                f"Roster has {roster['numeric_client_ids']} numeric Client IDs; store them as text."
            )
        if roster["duplicate_client_ids"]:
            result["warnings"].append(
                "Duplicate Client IDs block app lookup: "
                + ", ".join(roster["duplicate_client_ids"])
            )
        if roster["leavers_without_end_date"]:
            result["warnings"].append(
                f"Roster has {roster['leavers_without_end_date']} Leavers without an End date."
            )
    finally:
        workbook.close()

    result["ok"] = not result["errors"]
    return result


def main() -> int:
    parser = argparse.ArgumentParser(
        description="Check FTE Count.xlsx for the Power Platform PTO/Away contract."
    )
    parser.add_argument(
        "workbook", nargs="?", type=Path, default=Path("templates/FTE Count.xlsx")
    )
    parser.add_argument("--json", action="store_true", dest="as_json")
    args = parser.parse_args()
    result = inspect_fte_contract(args.workbook)
    if args.as_json:
        print(json.dumps(result, indent=2, default=str))
    else:
        print(f"Workbook: {result['path']}")
        print(f"Contract: {'PASS' if result['ok'] else 'FAIL'}")
        for error in result["errors"]:
            print(f"ERROR: {error}")
        for warning in result["warnings"]:
            print(f"WARNING: {warning}")
        roster = result["roster"]
        print(f"Roster populated rows: {roster['populated_rows']}")
    return 0 if result["ok"] else 1


if __name__ == "__main__":
    raise SystemExit(main())
